from pathlib import Path
import json,csv,zipfile,shutil,re,hashlib
import fitz,openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
R=Path('/workspace/project-review');O=R/'outputs/BRRP_P8_Input_Audit';rows=json.loads((O/'P8_file_inventory.json').read_text());pages=json.loads((O/'P8_page_index.json').read_text());ocr=json.loads((O/'P8_OCR_inventory.json').read_text())
uses={
'Addendum (1).pdf':('Amended vendor/agency and technical/commercial provisions','Apply with A9/A11/A16/B6 and payment/cost controls; earlier identical content retained'),
'Addendum.pdf':('Exact duplicate of Addendum (1).pdf','One canonical amendment content; retain both filenames in inventory'),
'Addendum_Structure_SOW.pdf':('ICG/ICP/BHS jacket strengthening before new bridge installation','Existing scope/design/MTO/installation predecessor; identical earlier copy'),
'AddendumtoA3.pdf':('Currency/import/local-content payment/pricing provisions','Commercial/bid-price/invoice currency controls; identical earlier copy'),
'Appendix A-9.pdf':('Bid technical confirmations and evidence requirements','OCR-reviewed; originalitem13 must be read with deleting addendum; no eligibility self-certification'),
'Appendix A-11 Offshore SVL Revision 12.1.pdf':('Suggested vendor reference list Rev12.1','Retain reference; 23Aprcover names12.3; current revision/applicability reconciliation open'),
'Appendix A-16 List of Agencies.pdf':('Activity-wise agency reference list Rev10.7','Read with reference-only amendment; qualifications/commitments remain to verify'),
'Appendix A-17.pdf':('Subcontractor/yard/marine commitment-letter formats','Prepare actual signed commitments; forms are not resource-availability evidence'),
'Appendix B-5.pdf':('Experience disclosure and supporting evidence','OCR-reviewed; reconcile actual BEC/amendments, consortium roles and threshold references'),
'Appendix B-6.pdf':('Work division, location and subcontracting restrictions','OCR-reviewed; retain bidder project management/required qualifying activity perBEC; no whole-works back-to-back subcontract'),
'Appendix C-8 DCI MCI.pdf':('Document/material control, predecessor stages, review and progress rules','Replace provisional review/progress assumptions; develop approved projectDCI/MCI'),
'Appendix C-9.pdf':('Prescribed vessel safety audit and critical pre-sail documents','Marine readiness/audit/certification/observation closure; actual audit not performed'),
'Appendix C-10.pdf':('Marine Safety Cell vessel/crew/equipment audit checklist','Vessel particulars, certificates, equipment service and crew evidence; actual clearance not evidenced'),
'Combined_Addendum_2.pdf':('Piping design/stress/fabrication/installation/insulation specifications','Piping stress/service-support/procurement/QA/test applicability; standards attached do not automatically add equipment scope'),
'Combined_Addendum_1 (1).pdf':('23Apr2026cover; engineering matrix; revised drawing index; actual structural and discipline sources','Close source-GA availability gaps; revise deliverables/DCI/schedule/specification mapping'),
'Combined_Addendum_1.pdf':('Exact duplicate of Combined_Addendum_1 (1).pdf','One canonical285page content; preserve second filename and hash linkage')}
findings=[
('F01','SOURCE GAP CLOSED','Combined_Addendum_1 (1).pdf','27','B774-101-83-43-0001 ICW-ICG lower-deck dimensioned GA/details is supplied.','Replace prior missing-GA status with SOURCE RECEIVED / TENDER / VERIFY; update design-input register.','Engineering / Document Control'),
('F02','SOURCE GAP CLOSED','Combined_Addendum_1 (1).pdf','31-32','B774-102-83-43-0001 ICP-ICG lower-deck two-sheet GA/details is supplied.','Register both sheets; develop model/member/connection input from source subject to survey and checking.','Engineering'),
('F03','SOURCE GAP CLOSED','Combined_Addendum_1 (1).pdf','36-41','BHS-SLQ new GA, landing modifications/extensions and demolition drawing are supplied.','Map actual103series to BHS installation/service-transfer/demolition engineering; source is not AFC approval.','Engineering / Offshore'),
('F04','ENGINEERING SCOPE EXPANDED/CLARIFIED','Combined_Addendum_1 (1).pdf','2','Annexure specifies10bridge analysis scopes;globaldeckinplace forICP/ICG/ICW/BHS/SLQ;localdeckinplace forWIS/NOQ/NQD/WIN/NC/SC1/SCA/SCF.','Create explicit calculation/deliverable/predecessor rows; verifyNOQ naming againstNQO; do not treat allplatformchecks as one generic activity.','Engineering / Planning'),
('F05','APPROVAL/RECORD CLASSIFICATION','Combined_Addendum_1 (1).pdf','2','Design basis/methods, reports and AFC drawings show approval;fabrication/as-built columns show records; matrix saysnotexhaustive.','Revise101itemdraftregister and authority/evidence fields; separate perbridge/landingmethodology submissions.','Engineering / QA / Document Control'),
('F06','REVIEW CALENDAR CORRECTION','Appendix C-8 DCI MCI.pdf','2-4;14-18','ProjectDCI/MCI must befinalised/approved;generally10workingdays review,listedcomplexfirstreviews longer,and subsequentreviews10workingdays.','Replace4calendardayplaceholder engineering review; build working-day calendars, reviewer capacity, staggered submissions and rework loops; recompute forecasts.','Planning / Engineering'),
('F07','PREDECESSOR STATE CONTROL','Appendix C-8 DCI MCI.pdf','3;17-18','DocumentcategoryA/I/R differs from predecessor stateS/R/A; out-of-sequence documents can be returned withoutreview.','Map mother documents and required predecessor status; enforce release conditions and preserve baseline revision history.','Document Control / Engineering'),
('F08','PROGRESS RULE CORRECTION','Appendix C-8 DCI MCI.pdf','4','ReviewCatC:noengineeringprogress;CatB:60%engineeringprogress onsubmission perwording;CatA:100%. MCI first/secondstage PS approval earns5%+5%of itemprocurementweight.','Reconcile/supersede earlierproposedinternal engineering/procurementcredit rules for contractual reporting; do not interpret as wholeprojectprogress or duplicate A-10billing.','Project Controls / Contracts'),
('F09','PRE-ACTIVITY SUBMISSION/MWS','Appendix C-8 DCI MCI.pdf','17-18;otherapplicablerows','Starred notes call for30dayadvance submission;double-starred items requireMWS approval whereapplicable.','Add submission/assurance milestones and review buffers to relevant operation logic; check row applicability,notblanketassumeallitems.','Marine / Planning'),
('F10','MARINE AUDIT GATES','Appendix C-9.pdf;Appendix C-10.pdf','C9pp1-21 especially20;C10pp1-19','Vessel particulars, POB, class/statutory/crew/equipment evidence and criticalpre-sail documents are prescribed.','Create candidate-vessel audit dossiers; schedule inspections/observationclosure/clearances before mobilisation or sail-away as applicable.','Marine / HSE / Assurance'),
('F11','SUBCONTRACTING CONSTRAINT','Appendix B-6.pdf;Appendix A-9.pdf','B6p1;A9pp4-7','Work division/locations/commitments and retainedbidderPM/qualifyingmainactivity requirements apply; wholeworksback-to-backsubcontract notallowed.','Map retained and subcontracted scope against actualBEC route; validate bidder/consortium/subcontractor evidence before commitments.','Contracts / Project Director'),
('F12','VENDOR REVISION CHECK OPEN','Combined_Addendum_1 (1).pdf;Appendix A-11 Offshore SVL Revision 12.1.pdf','Coverp1;A11pp1-3','CoverreferencesSVL12.3;separately uploadedlist is12.1. Generaladdendum makeslists reference-only subjecttechnical/tenderrequirements.','Do notlabel12.1latest; confirm12.3/currentapplicability while retainingqualification/exclusion and technical controls.','Procurement / Contracts'),
('F13','REVISED SCC/A3 ATTACHMENT MATCH OPEN','Combined_Addendum_1 (1).pdf','Coverp1','Coverannounces revisedSCC/A3; corresponding full revised attachments are not positivelyidentified inside thetwo combined PDFs.','Match against earlier/later separatelysuppliedfiles and officialrevision evidence; filename/metadataalone is not precedence proof. Do not freeze commercial/date basis from cover alone.','Contracts / Document Control'),
('F14','DRAWING REGISTER RECONCILIATION','Combined_Addendum_1 (1).pdf','3-5;27','Revisedregister omits some olderstandalone0005references;101GA registershowsRev0 while sheettitleblockshowsA and revisionhistory0.','Resolve sheet/revision identity; mapold/newreferences; do not keep requesting omittedold0005sheets as definite gaps or assume workdeleted.','Document Control / Engineering'),
('F15','DISCIPLINE COMPLETENESS','Combined_Addendum_1 (1).pdf;Combined_Addendum_2.pdf','C1pp42-285;C2pp1-250','Substantial electrical, instrumentation and piping specifications/ITPs/standards are now available.','Build scope-specific specification/MTO/vendor/test matrix; piping stress reactions must feed structure; attached genericstandards do not automatically create additional facilities.','Discipline Leads / Procurement / QA'),
('F16','FORM/VERSION COMPLIANCE','Appendix A-9.pdf;Addendum (1).pdf;Appendix B-5.pdf','A9pp4-5;Addendump1;B5p1','OriginalA9item13 yard-information provision is deleted byaddendum;experienceforms and qualifyingthresholdreferences requireBEC cross-check.','Maintain amendedcompliance matrix; do notapplydeletedclause ordeclareeligibilityfromaformalone.','Contracts / Bid Compliance')]
# Document boundaries for discipline applicability (index is not a claim of code-by-code engineering validation).
bound=[]
for name,pp in pages.items():
 for x in pp:
  z=x['text']
  if re.search(r'(?:PAGE|Page|page|SHEET|Sheet)\s*:?\s*1\s*(?:OF|of|/)',z) and (name.startswith('Combined') or name.startswith('Appendix C-8')):
   bound.append([name,x['page'],z[:450],'INDEXED; specific scope applicability to be confirmed'])
# Actual drawing locations, not only index references.
draw=[
('B774-000-83-43-0001','Generalnotes',6,7,'Reference source;register2sheets'),
('B774-101-83-43-0001','ICW-ICG LD GA/details',27,27,'Tender;revisionconflict register0/titleA'),
('B774-101-83-43-0002','ICWlandingmodification',28,28,'Tender;verifycurrentrevision'),
('B774-101-83-43-0003','ICWlandingextension',29,29,'Tender;verifycurrentrevision'),
('B774-101-83-43-0004','ICGlandingmodification',30,30,'Tender;mapassociatedextensiondetails'),
('B774-102-83-43-0001','ICP-ICG LD GA/details sheets1and2',31,32,'Tender;RevA inrevisedregister'),
('B774-102-83-43-0002','ICPlandingmodification',33,33,'Tender'),
('B774-102-83-43-0003','ICPlandingextension',34,34,'Tender'),
('B774-102-83-43-0004','ICGlandingmodification',35,35,'Tender'),
('B774-103-83-43-0001','BHS-SLQ new GA/details',36,36,'Tender;notAFC'),
('B774-103-83-43-0002','BHSlandingmodification',37,37,'Tender'),
('B774-103-83-43-0003','BHSlandingextension',38,38,'Tender'),
('B774-103-83-43-0004','SLQlandingmodification',39,39,'Tender'),
('B774-103-83-43-0005','SLQlandingextension',40,40,'Tender'),
('B774-103-83-43-0006','BHS-SLQ existingbridge demolition',41,41,'Tender;engineering/operationalrelease stillrequired')]
applications=[]
for r in rows:
 topic,use=uses[r['name']]
 method='Text extraction + page index'
 if r['low_text_pages']:method+='; OCR of low-text pages'
 applications.append([r['name'],r['pages'],r['bytes'],r['sha256'],r['duplicate_in_batch'] or '',len(r['prior_matches']),method,topic,use,'Registered/version-checked; relevant clauses/source drawings reviewed; full design/compliance validation not claimed'])
# Workbook
w=openpyxl.Workbook();w.remove(w.active)
def tab(name,head,data,widths=None):
 s=w.create_sheet(name);s.append(head)
 for r in data:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
 for j in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(j)].width=widths[j-1] if widths and len(widths)>=j else 45
 for rr in s:
  for c in rr:c.alignment=Alignment(wrap_text=True,vertical='top')
 return s
tab('Batch Scope',['Item','Evidence'],[['Uploadedfiles',16],['Uniquecontentswithinbatch',14],['UniquePDFpageswithinbatch',754],['Newuniquecontentsversusearlierlocalcopies',6],['Exactsamebatchduplicategroups','Addendum.pdf=Addendum(1).pdf;Combined_Addendum_1.pdf=Combined_Addendum_1(1).pdf'],['OCRscreenedlowtextpages',len(ocr)],['Auditlimit','This register covers the latest16uploads. Allhistoricalupload/fileincorporationaudit remainsseparate. No fullindependentdesign/eligibility/contractvalidation claimed.']],[45,140])
tab('File Incorporation',['Filename','Pages','Bytes','SHA256','Samebatchduplicateof','Earlierexactmatches','Extraction/reviewmethod','Content','Projectapplication','Reviewstatus'],applications,[65,12,18,72,65,20,50,95,120,100]);tab('Findings and Changes',['ID','Classification','File','PDFpages','Verifiedfinding','Requiredpackagechange','Owner'],findings,[14,45,95,45,125,125,45]);tab('Structural Drawing Locations',['DrawingID','Description','FirstPDFpage','LastPDFpage','Status'],draw,[45,65,23,23,110]);tab('Discipline Page Index',['File','Startpage','Extractedheading','Reviewlimit'],bound,[65,20,140,80]);tab('OCR Coverage',['File','Page','OCRcharacters','Engineexit','OCRpath','OCRsnippet'],[[x['name'],x['page'],x['chars'],x['exit'],x['ocr_path'],x['snippet']] for x in ocr],[70,16,20,20,110,140]);
w.save(O/'BRRP_P8_Input_Inventory_and_Impact_Register.xlsx')
# Report.
d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10);d.add_heading('BRRP — Latest Upload Batch Analysis and Incorporation Audit',0)
d.add_paragraph('P8 |25October2026 | Latest16uploadedfiles | Not a completed all-history audit or design approval')
d.add_paragraph('All16files were downloaded, hashed, compared to earlier local copies and registered. They contain14uniquecontents and754pages withinthisbatch. Sixuniquecontents were not found among earlierdownloadedcopies. Twoexactduplicatepairs are preserved byfilename but analysed throughcanonicalcontent. Textwasextracted/indexedforuniquePDFs;20low-textpages wereOCR-screened; criticaldrawingandclauseevidence was reviewed. This doesnotclaim an independentfulltechnicalcheckofeverypage or everyhistoricalupload.')
def T(head,data):
 t=d.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,head):c.text=str(v)
 for row in data:
  for c,v in zip(t.add_row().cells,row):
   c.text=str(v)
   for p in c.paragraphs:
    for run in p.runs:run.font.size=Pt(8)
 return t
d.add_heading('1. Principal conclusions',1)
for t in ['The main dimensioned tender GAs for allthree replacement/newbridges are now positivelylocated, together with two-sheetICPdetails andBHSlanding/demolitiondrawings. Earlier source-availability gaps must be updated; finalverification/AFCstatus is not implied.','The23April2026addendum supplies a specific bridge/platformengineeringmatrix. Fiveplatforms requireglobaldeckinplaceanalysis and eightlocaldeckinplaceanalysis in theannexure; thelist isnotexhaustive.','AppendixC8provides actualDCI/MCIreview, predecessor, engineeringprogress and two-stagePSrequirements. The earlier4calendardayreviewplaceholder and proposedprogressrules require correction/reconciliation.','C9/C10provide vessel-audit and criticalpre-sail requirements; blankforms are not clearanceevidence.','B6/A9specify retainedbidderPM/qualifyingactivity, workdivision andcommitmentevidence; qualifications cannotbeinferredfromcapabilityclaims.','Thecoverreferences revisedSCC/A3andSVL12.3, but their actualrevisedattachments are notpositivelymatched in thesetwocombinedPDFs. The separatelyuploadedSVLis12.1. Resolveversionlinks beforefreezingcommercial/vendorbasis.']:
 d.add_paragraph(t,style='List Bullet')
d.add_heading('2. File-by-file incorporation',1)
T(['File','Pages','Content / use','Version/duplicate'],[[r['name'],r['pages'],uses[r['name']][0]+'; '+uses[r['name']][1],('Samebatchduplicate: '+r['duplicate_in_batch']) if r['duplicate_in_batch'] else ('Exactearliercopy' if r['prior_matches'] else 'Newuniquecontent')] for r in rows])
d.add_heading('3. Changes to engineering source status',1)
T(['Drawing','Description','CombinedAddendum1pages','Status'],[[x[0],x[1],str(x[2])+'-'+str(x[3]),x[4]] for x in draw])
d.add_paragraph('Theactualsheetsremainissuedfortenderandexpresslyrequirecontractorverification. Thesenowprovideconcretegeometry/member/connection references fordesign development; they are not signedAFCdrawings. PreserveoriginalsourceandproposednewdrawingIDs separately. Theupdateddrawingregistermustbe reconciledwithsheetrevisionblocks andoldscope references; omissionfromanindex alone doesnotdeletecontractwork.')
d.add_heading('4. Specific findings and package impacts',1)
for id,kind,file,pp,finding,action,owner in findings:
 d.add_heading(id+' — '+kind,2);d.add_paragraph('Source: '+file+';PDFpages '+pp+'.');d.add_paragraph(finding);d.add_paragraph('Requiredchange: '+action);d.add_paragraph('Owner: '+owner+'. Implementation/approvalstatus: OPEN.')
d.add_heading('5. DCI/MCI rules to implement',1)
for t in ['Developandobtainapprovalforproject-specificDCI/MCI;standardlistisnotexhaustive. Useproject/facility/discipline/document-type numberingandlinkthetemporaryShivanshIDs.','KeepdocumentcategoryA/I/R(approval/information/review) distinctfromreviewstatusCatA/B/C. Keepmother-documentstateS/R/Adistinctfromboth.','Generallyallow10workingdaysforreview;uselongerfirst-reviewtimeswhereexplicitlylisted;subsequentrevisions10workingdays. Useanagreedworkingcalendar,not10calendardaysbyassumption.','Engineeringcredit:CatCnone,CatB60%perwording,CatA100%. Applyapproveddocumentweights;theseare notwhole-projectcompletionpercentages. Earlierproposedinternalmilestoneweightscannotreplaceapplicablecontractualrules.','Purchase-specificationfirst/secondstageapprovalhas5%+5%ofitemprocurementweightperC8;donotapplyas10%ofthewholeproject orduplicatepaymentmilestones.','Submitrevisionswithcommentresolutionsheetandclouded/indexedchanges;endorsethroughcontractordesignagency. Out-of-sequencedocumentsmaybereturnedwithoutreview.','Apply30dayadvance-submissionandMWSrequirementswhereidentifiedbythetablenotes and relevantrow;notasablanketrequirementforallgenericfacilities.']:
 d.add_paragraph(t,style='List Bullet')
d.add_heading('6. Schedule and commercial status',1)
d.add_paragraph('Do not reaffirm the15June2027scenario asvalidatedafterreceipt oftheseinputs. Addbridge/platformanalyses,separatereviews,resubmissions,PSstages,vesselauditsandapplicableadvanceassurancemilestones, then recomputethelogic/resources/calendar. Theoldscenario remains historical/preliminary until revised. Procurementvendor/agencyrulesmustbereadwiththegeneraladdendumthatmakesthelistreference-only;thisdoesnotbyitselfestablishqualification,suitability,availabilityorpermissiontouseexcludedfirms.')
d.add_heading('7. Outstanding exact-version and interpretation items',1)
for t in ['MatchrevisedTechnicalSCCandA3announcedonthe23Aprilcover to actualcurrentattachments;metadataandfilenamesalone are insufficient.','ConfirmSVL12.3orcurrentapplicability; do notcalltheuploaded12.1listlatest.','ReconcileGA101registerRev0,titleblockRevAandhistoryRev0.','Mapold101/1020005references to revisedregister/sheetcontent; do notassumeworkdeletedorrequestobsoleteIDswithoutchecking.','ConfirmNOQlabelinannexureagainstNQOprojectnaming.','Reconcileexperience/formthresholdreferences againstactualBECandamendments; do notdeclarebiddereligibilityfromformsalone.','Engineeringcalculation,as-builtsurveyverification,actualsupplier/crew/vesselcommitments,AFCapprovalandphysicalexecutionremainseparateopenitems.']:
 d.add_paragraph(t,style='List Bullet')
d.add_heading('8. Scope of confirmation',1)
d.add_paragraph('Confirmed:everyfileinthislatest16-filebatchisregisteredandversion-checked,andkeyrelevantrequirements/sourceavailabilitychangesareidentifiedwithpagereferences. Notconfirmed:allhistoricaluploadsfullyincorporated;everygenericstandard/itemapplicable;allclausesresolved;orproject/design/AFCcompletion. Acompleteall-historyauditmustretainthesedistinctions.')
d.save(O/'BRRP_P8_Latest_Upload_Analysis_and_Required_Updates.docx')
# Source extract preserves actual sheets. It is not a new engineering drawing issue.
row=next(x for x in rows if x['name']=='Combined_Addendum_1 (1).pdf');src=fitz.open(row['path']);ext=fitz.open()
for n in [1,2,3,4,5,6,7]+list(range(27,42)):ext.insert_pdf(src,from_page=n-1,to_page=n-1)
ext.save(O/'P8_Source_Structural_Drawings_and_Deliverable_Matrix.pdf')
# Explicit source progress / review rules machine-readable for later schedule regeneration.
(O/'P8_source_rules.json').write_text(json.dumps({'general_review_working_days':10,'review_status_engineering_credit':{'Cat C':0,'Cat B':60,'Cat A':100},'PS_stage_item_weight_credit_percent':[5,5],'scope':'SourceC8rules subjectprojectDCI/MCIapplicability;notglobalprojectprogress','global_deck_platforms':['ICP','ICG','ICW','BHS','SLQ'],'local_deck_platform_labels_as_printed':['WIS','NOQ','NQD','WIN','NC','SC1','SCA','SCF'],'drawing_source_locations':draw},indent=2))
with open(O/'P8_File_Incorporation.csv','w',newline='') as f:
 z=csv.writer(f);z.writerow(['Filename','Pages','Bytes','SHA256','Duplicateof','Earliermatches','Method','Content','Application','Reviewstatus']);z.writerows(applications)
(O/'P8_findings.json').write_text(json.dumps(findings,indent=2))
print('Prepared16-fileinventory;16specificfindings;15drawing-locationrecords;6-sheetworkbook;Wordimpactreview;22-pagesourceextract. Fullhistory/designvalidationnotclaimed.')
