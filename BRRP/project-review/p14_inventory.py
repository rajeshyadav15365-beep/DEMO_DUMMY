import pathlib,json,hashlib,zipfile,fitz,openpyxl
from docx import Document
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P14_Consolidated_Management');rs=json.loads(pathlib.Path('/workspace/project-review/p14_manifest.json').read_text());curr={r['path']for r in rs if 'path'in r};prior={}
for f in pathlib.Path('/workspace/attachments').rglob('*'):
 if f.is_file()and str(f)not in curr and f.suffix.lower()in ['.pdf','.docx','.xlsx','.csv','.zip','.xml']:prior.setdefault(hashlib.sha256(f.read_bytes()).hexdigest(),[]).append(str(f))
seen={};es=[]
for r in rs:
 if 'path'not in r:continue
 f=pathlib.Path(r['path']);h=hashlib.sha256(f.read_bytes()).hexdigest();r.update(name=f.name,sha256=h,prior=prior.get(h,[]),duplicate=seen.get(h));seen.setdefault(h,f.name)
 if r['duplicate']:continue
 e={'name':f.name,'path':str(f)}
 if f.suffix=='.xlsx':
  w=openpyxl.load_workbook(f,data_only=False);e['sheets']=[]
  for s in w:e['sheets'].append({'name':s.title,'rows':s.max_row,'cols':s.max_column,'cells':{c.coordinate:str(c.value)for row in s for c in row if c.value is not None}})
 elif f.suffix=='.docx':
  d=Document(f);e['paragraphs']=[q.text for q in d.paragraphs if q.text.strip()];e['tables']=[[[c.text for c in row.cells]for row in t.rows]for t in d.tables]
 elif f.suffix=='.pdf':
  d=fitz.open(f);r['pages']=len(d)
  if not r['prior']:e['pages']=[{'page':i+1,'text':' '.join(pg.get_text().split())}for i,pg in enumerate(d)]
 elif f.suffix=='.zip':
  with zipfile.ZipFile(f)as z:e['archive_members']=[{'name':n,'sha256':hashlib.sha256(z.read(n)).hexdigest()}for n in z.namelist()if not n.endswith('/')];e['archive_integrity']=z.testzip()
 es.append(e)
(p/'P14_latest_inventory.json').write_text(json.dumps(rs,indent=2));(p/'P14_latest_extracts.json').write_text(json.dumps(es,indent=2))
print(json.dumps({'supplied':len(rs),'opened':sum('path'in r for r in rs),'distinct':len(seen),'new_distinct':len({r['sha256']for r in rs if 'sha256'in r and not r['prior']}),'new':[(r['name'],r.get('pages'))for r in rs if 'sha256'in r and not r['prior']and not r['duplicate']]},indent=2))
