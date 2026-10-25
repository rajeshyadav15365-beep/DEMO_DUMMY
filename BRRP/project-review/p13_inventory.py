import pathlib,json,hashlib,fitz
from docx import Document
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P13_Master_Plan_Review');rows=json.loads(pathlib.Path('/workspace/project-review/p13_manifest.json').read_text());curr={r['path']for r in rows if 'path'in r};prior={}
for f in pathlib.Path('/workspace/attachments').rglob('*'):
 if f.is_file() and str(f)not in curr and f.suffix.lower()in ['.pdf','.docx','.xml','.xlsx']:
  prior.setdefault(hashlib.sha256(f.read_bytes()).hexdigest(),[]).append(str(f))
seen={};extracts=[]
for r in rows:
 if 'path'not in r:continue
 f=pathlib.Path(r['path']);h=hashlib.sha256(f.read_bytes()).hexdigest();r.update(name=f.name,sha256=h,prior_matches=prior.get(h,[]),batch_duplicate_of=seen.get(h));seen.setdefault(h,f.name)
 if r['batch_duplicate_of']:continue
 e={'name':f.name,'path':str(f)}
 if f.suffix=='.docx':
  d=Document(f);e['paragraphs']=[{'paragraph':i+1,'text':q.text}for i,q in enumerate(d.paragraphs)if q.text.strip()];e['tables']=[[[c.text for c in row.cells]for row in t.rows]for t in d.tables]
 elif f.suffix=='.pdf':
  d=fitz.open(f);r['pages']=len(d)
  if not r['prior_matches']:e['pages']=[{'page':i+1,'text':' '.join(q.get_text().split())}for i,q in enumerate(d)]
 extracts.append(e)
(p/'inventory.json').write_text(json.dumps(rows,indent=2));(p/'unique_content_extracts.json').write_text(json.dumps(extracts,indent=2))
print(json.dumps({'uploads':len(rows),'successful':sum('path'in r for r in rows),'distinct':len(seen),'new_distinct':len({r['sha256']for r in rows if 'sha256'in r and not r['prior_matches']}),'unique_groups':[{'name':r['name'],'prior':len(r['prior_matches']),'pages':r.get('pages')}for r in rows if 'sha256'in r and not r['batch_duplicate_of']]},indent=2))
for e in extracts:
 if e['name']=='ticket.pdf':print('TICKET classification excerpt:',str(e.get('pages'))[:1000])
