from pathlib import Path
import ast,re,json,zipfile,shutil,hashlib
from docx import Document
import openpyxl
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.pagesizes import A4
from reportlab.lib.colors import HexColor
R=Path('/workspace/project-review');O=R/'outputs/BRRP_P7_Assurance_Completion'
# Reuse only the pure text-cleaning definitions, never rerun older publication scripts.
module=ast.parse((R/'beautify_p5.py').read_text());nodes=[]
for node in module.body:
 if isinstance(node,ast.Assign) and any(isinstance(t,ast.Name) and t.id in ['pairs','mp'] for t in node.targets):nodes.append(node)
 if isinstance(node,ast.FunctionDef) and node.name=='clean':nodes.append(node)
ns={'re':re};exec(compile(ast.Module(body=nodes,type_ignores=[]),'text_cleanup','exec'),ns);clean=ns['clean']
# Replace dense draft narrative with properly written final prose.
replacements={
'A101-item':'A register containing101engineering deliverables is not evidence that101designs or drawings have been completed. The nineP4sheets comprise source-informed arrangements, functional support interfaces and detailing schedules. They are not nine fully engineered drawings. The package is not ready for baseline approval while design, cost, resource, procurement, weather and review evidence remain unresolved.',
'The existing205':'The205activity scenario has one start and one terminal milestone, and its dependency, calendar and named-resource sequencing checks passed. Those checks establish internal consistency only. The25procurement packages do not have individual vendor-backed lead times. P010 is a generic25day bulk-material workflow. Crew, yard, inspection, coating and POB capacities and most duration inputs are unverified. The10day weather hold is not an operation-specific weather forecast. The15June2027target remains an illustrative scenario until these gaps are resolved.',
'The earlier3day':'The earlier three-day material-delay test measures the sensitivity of the coarse seasonal calendar and reserve assumptions. It does not prove that actual demobilisation must wait until October. Final planning requires actual operating limits, workability, safe suspension, restart and demobilisation strategies. Contractual relief and cost entitlement are separate assessments.',
'Closure: project-specific':'Closure requires a project-specific procedure or populated register, the designated review, objective evidence and updates to affected baseline documents. Drafting a procedure does not establish its implementation or approval.',
'Proposedcode:':'Use a common code such as BRRP/<bridge>/<workstream>. ENG,MAT,FAB,MAR,OFF and COM identify engineering, materials, fabrication, marine, offshore construction and completions. Each control account requires one accountable manager, scope, budget/hours, deliverables, schedule and measurement rules. Purchase items, material issues, timesheets, accepted quantities, invoices and turnover records must use compatible codes. Allocate shared costs once.',
'The following are proposed internal physical-progress':'The following are proposed internal physical-progress rules, separate from contractual billing milestones. Approve them before use and align them with accepted work and budgeted hours/cost. Do not assign actual progress without evidence. Discipline weights total100percent; overall project weighting needs an approved basis. Drawing counts, steel tonnage, services and cash cannot be averaged as one physical unit.',
'Weightedprogress=':'Weighted progress equals the sum of approved milestone weight multiplied by evidenced completion fraction. Use consistent measurement and acceptance rules and preserve the approved baseline. Earned value(EV) is the approved budgeted value of performed work; actual cost(AC) covers the same scope, period and currency; planned value(PV) is the approved time-phased plan. CPI=EV/AC and SPI=EV/PV require valid compatible inputs. No project CPI or SPI is calculated here because those inputs are not established.',
'Fixa singledata':'Fix one reporting data date and named data owners. Collect evidenced actual starts/finishes, accepted quantities, inspection status, remaining work and actual resource/cost records.',
'Validateactuals':'Validate actuals and resolve out-of-sequence work. Preserve the baseline and approved change history. Forecast dates must not be forced to match a target.',
'Recalculatecritical/near-criticalpaths':'Recalculate critical and near-critical paths, forecast cost/cash, weather exposure and procurement/resource constraints.',
'Escalate when a critical prerequisite':'Escalate a critical prerequisite missing its required date, an operational gate on HOLD, an uncommitted supplier window or an uncovered funding requirement. Approve internal response times separately; do not represent them as contractual limits.',
'Issue weeklyreport':'Issue a weekly report identifying each decision, accountable owner, need-by date, evidence and consequence. Build the six-week look-ahead from executable work fronts.',
'The appendedworkbookmaps':'The workbook maps all25procurement packages to need points, leaving unknown supplier lead times and dates blank. Obtain vendor-document, approval, manufacture, inspection and delivery milestones; connect them to genuine work-front needs. Validate shared inspection/coating, support vessels, transfers and POB as well as fabrication lines and crews. No purchase, charter or external contact has been performed.',
'Agree clientterminology':'Agree the client terminology and certificate criteria for construction completion, readiness for testing, readiness for service and contractual acceptance. These are proposed planning categories, not asserted contract terms. Define test values, punch classification, reinstatement/isolation, applicable manuals/spares/training and asset-data requirements. BHS demolition remains dependent on new-bridge/service acceptance and engineered removal release.',
'Every applicablecontract':'Every applicable contract, scope and design requirement must map to a deliverable, verification method, owner and acceptance authority.',
'Scope,MTO,engineering':'Scope, MTO, engineering, procurement, schedule, resources, cost/cash, methods, quality/HSE and turnover must be internally consistent.',
'Project-specificinputs':'Resolve project-specific inputs and uncertainty, or obtain an authorised decision on a permissible residual risk. A management exception must not conceal missing technical adequacy.',
'Reviewfindings':'Close review findings against actual evidence. Perform suitably independent reviews where required.',
'Issuedfiles':'Check issued-file readability and successful recipient receipt/opening through the agreed delivery route. Local checksums do not prove delivery.',
'Actualengineeringapproval':'State engineering approval, AFC release, supplier commitments, physical work and client acceptance only when supported by evidence.',
'Nextstep:':'Next steps are to populate project-specific inputs, resolve the priority engineering/schedule/commercial gaps and complete competent reviews. This self-review improves the preparation but does not certify completeness or establish superiority over another EPC company.'}
p=O/'BRRP_P7_Management_Package_Review_and_Remaining_Details.docx';doc=Document(p)
for para in doc.paragraphs:
 t=para.text
 for prefix,new in replacements.items():
  if t.startswith(prefix):t=new;break
 para.text=clean(t)
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for para in cell.paragraphs:para.text=clean(para.text)
from docx.shared import Pt
for table in doc.tables:
 for row in table.rows:
  for cell in row.cells:
   for para in cell.paragraphs:
    for run in para.runs:run.font.size=Pt(8)
doc.save(p)
p=O/'BRRP_P7_Package_Assurance_and_Closure.xlsx';w=openpyxl.load_workbook(p)
for s in w:
 for row in s:
  for c in row:c.value=clean(c.value)
w.save(p)
# Review PDF: all24findings and their four closure actions, with actual maturity stated.
j=json.loads((O/'P7_review_data.json').read_text());styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='Compact',fontSize=9,leading=12,spaceAfter=5));story=[]
def P(t,sty='Compact'):story.append(Paragraph(clean(t).replace('&','&amp;'),styles[sty]));story.append(Spacer(1,4))
P('BRRP — Management Package Review and Completion Supplement P7','Title')
P('25 October2026 | Self-review | Preliminary package; no superiority or independent-assurance claim')
P('Review conclusion','Heading1')
P('The current package cannot be substantiated as better than leading offshore EPC companies. Its scope/workstream preparation is useful, but structural design, supplier/resource evidence, schedule/weather validation, cost funding and real approvals are incomplete. Document and activity counts do not prove readiness.')
P('Material corrections','Heading1')
for x in ['The205activitycriticalpath is a calculation for an assumed network, not a validated criticalpath for the whole actual project.','All25procurementpackages require individual leadtime/approval/delivery logic where relevant; a single bulk-materialworkflow is insufficient.','Replace the10dayweatherhold with operation-specific workability, interruption, safe suspension and demobilisation analysis.','A101itemregister and nine development sheets are not101completed designs or nine construction drawings.','File integrity is verified locally, but recipient download/opening is still not confirmed.']:P(x)
for id,title,kind,finding,owner,gate,actions in j['findings']:
 P(id+' — '+title,'Heading2');P('Classification: '+kind+'. '+finding+' Owner: '+owner+'. Need-by: '+gate+'.')
 for k,a in enumerate(actions,1):P(str(k)+'. '+a)
P('Acceptance of this supplement','Heading1');P('P7adds24controls and96closureactions to P6; it does not prove those actions implemented. It includes18starter requirement links,60controlaccountmappings,25procurementneed links, proposedprogressrules, distinctchange workflows and7turnoverboundaries. Project-specific values, evidence, competent reviews and required approvals remain prerequisites.')
SimpleDocTemplate(str(O/'BRRP_P7_Package_Review_Supplement.pdf'),pagesize=A4,leftMargin=40,rightMargin=40,topMargin=36,bottomMargin=36).build(story)
# Validate data, traceability and no invented readiness/progress/leadtime.
assert len(j['findings'])==24 and len(j['actions'])==96
assert len(j['requirements'])==18 and len(j['control_accounts'])==60 and len(j['procurement_links'])==25
for family in set(x[0] for x in j['progress']):assert sum(x[2] for x in j['progress'] if x[0]==family)==100
assert all(x[3] is None and x[4] is None and x[5] is None for x in j['procurement_links'])
assert len(w.sheetnames)==10
for s in w:
 assert not any(c.data_type=='e' for row in s for c in row)
for p in O.iterdir():
 if p.suffix in ['.docx','.xlsx']:
  with zipfile.ZipFile(p) as z:assert z.testzip() is None
(O/'README_P7.txt').write_text('P7 is an evidence-based self-review and management-package completion supplement. It is NOT an independent technical review, certification, benchmark proof, AFC approval or completed structural design. It qualifies the P5schedule and delivery-readiness claims.24findings/96closureactions supplement104P6actions. Unknownvendorleadtime and actualprogress remainunassigned.\n')
manifest=[]
for p in sorted(O.iterdir()):
 if p.is_file():manifest.append({'file':p.name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
(O/'MANIFEST_P7.json').write_text(json.dumps(manifest,indent=2))
Z=O.parent/'BRRP_P7_Review_and_Completion_Supplement.zip'
with zipfile.ZipFile(Z,'w',zipfile.ZIP_DEFLATED) as z:
 for p in O.iterdir():z.write(p,arcname=O.name+'/'+p.name)
with zipfile.ZipFile(Z) as z:assert z.testzip() is None
for p in [Z,O/'BRRP_P7_Management_Package_Review_and_Remaining_Details.docx',O/'BRRP_P7_Package_Assurance_and_Closure.xlsx',O/'BRRP_P7_Package_Review_Supplement.pdf']:
 q=Path('/workspace/downloads')/p.name;shutil.copyfile(p,q);q.chmod(0o644)
print('P7validated:24controls/96actions,18requirementlinks,60controlaccounts,25unknownleadtime-preservingprocurementlinks,5progressfamiliesweights100%,10workbooksheets. Files saved in downloads. PDF pages:',end=' ')
import fitz
print(len(fitz.open(O/'BRRP_P7_Package_Review_Supplement.pdf')))
