import pathlib,json,shutil
from docx import Document
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,Alignment
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P13_Master_Plan_Review');inv=json.loads((p/'inventory.json').read_text());a=json.loads((p/'XML_schedule_audit.json').read_text());links=a['link_records'];pred={x['successor']for x in links};succ={x['predecessor']for x in links};tasks=a['tasks'];unique={r['sha256']for r in inv if 'sha256'in r};new={r['sha256']for r in inv if 'sha256'in r and not r['prior_matches']}
findings=[
('Master-plan status','Volumes4–6 document status and source-basis paragraphs','Volumes describe pre-award management planning, not construction release. Repeated copies are not proof of implementation or approval.','Use these as management frameworks; link every bridge-specific deliverable to a controlled drawing/calculation, owner, due date and acceptance record.'),
('Marine plan','Volume4 clauses4.1–4.10','Four proposed campaigns are management groupings, not four proven charters. Vessel selection, engineered module weights/COG, transport/lift/removal checks, permits and go/no-go gates remain evidence requirements.','Close module-by-module geometry/weight/COG and rigging; crane radius/configuration, deck strength/stability, positioning, weather workability, assurance and charter availability. Never treat282MT BHS package quantity as a single lift weight.'),
('Internal logistics targets','Volume4 clause4.6','14-day cargo readiness and72-hour manifest are proposed management controls. They are not established here as ONGC contractual notice periods.','Mark proposed internal targets separately from mandatory contractual notices; coordinate actual shore-base and vessel booking requirements.'),
('Schedule planning windows','Volume5 clauses5.1,5.3,5.5,5.10','Provisional day-number windows and360-day target are expressly unvalidated. Volume5 states feasibility is not demonstrated.','Reconcile12months-from-NOA SCC against the360-day planning target; use actual commencement evidence, review-working-day calendars, monsoon conditions and resource capacity before approving baseline.'),
('XML structural checks','ProjectLibre XML Tasks/PredecessorLink; XML_schedule_audit.json','207 non-summary tasks,124 predecessor links. No duplicate task UIDs, missing referenced predecessor IDs or directed cycles found by structural screening. All links have Type1 (FS in Microsoft Project XML).','These checks establish parseable link structure only. Import compatibility and scheduling-engine calculations were not tested in ProjectLibre/P6.'),
('Open network ends','XML link audit','83 tasks have no predecessor;83 have no successor;17 are isolated. Multiple package starts/finishes can be intentional, but no integrated completion-path proof follows.','Review every open end and connect required engineering, procurement, inspection, fabrication, logistics, marine, installation and turnover gates. Add genuine project start/completion logic; do not add arbitrary links merely to reduce counts.'),
('XML calendar and duration semantics','XML StartDate/Calendars/MinutesPerDay; task Notes','Start is assumed1November2026 at08:00; calendar is Standard5-day Placeholder with8hours/day and no calendar exceptions. Example first survey duration224hours corresponds28workingdays; provisional day windows in Notes are not calculated dates.','Establish calendar meaning for each task. Separate working-day engineering review, yard shifts, marine offshore calendars, weather/season restrictions and approved relief; re-evaluate durations from quantities/productivity.'),
('Missing resource and CPM evidence','XML Resources/Assignments and task fields','No resources or assignments are present. Examined task fields do not establish calculated start/finish, total float or an accepted critical path.','Assign named capacity/units, work shifts, vessels, crews and workfront constraints. Run CPM then resource feasibility and sensitivity checks; retain scheduler settings and calculation logs. No revalidated finish date is claimed by this review.'),
('Commercial and progress controls','Volume6 clauses6.1–6.4,6.9','Cost/commitment/accrual/actual, earned value, certified progress and cash receipts are distinct controls. Generic framework text does not populate project budgets or demonstrate bank funding.','Link approved WBS quantities/rates and cost accounts to activity progress rules, amended A-10 milestones, invoicing/e-MB and full-life cash flow. Apply P9/P11 security, tax and amendment findings.'),
('Quality, HSE and turnover','Volume6 clauses6.5–6.10','QA/HSE/emergency/risk/turnover provisions specify required systems and evidence; tables and narrative do not establish that tests, drills, inspections or handovers occurred.','Issue scoped ITP/WPS/PQR/NDT/coating matrices; competent personnel and equipment records; HIRA/JSA/PTW/SIMOPS and emergency interfaces; bridge-wise completion, service reinstatement, scrap and acceptance dossiers.'),
('Benchmark claims','Volumes4–6 external EPC benchmark statement','References to Saipem and McDermott are general benchmarks. This batch supplies no independent comparative performance evidence or verified internal company manuals.','Describe adopted practices and measurable controls without claiming superiority to either EPC contractor.'),
('Reconciliation carry-forward','Volumes4–6 quantity basis; P8–P12 reviews','Repeated SOW/specification/tender material does not close MTO discrepancies, design questions, revised SCC/A-3 issue identity or engineering approvals.','Maintain open-item history, source precedence and amendment overlays; do not mark rows closed because a file was re-uploaded.'),
('Unclassified file','ticket.pdf; unique-content extract','Classify separately based on actual content. A personal/travel document does not change project technical or commercial scope merely because uploaded with tender documents.','Exclude unrelated personal information from shared project deliverables; do not treat a ticket as project award, approval or mobilisation evidence without a relevant explicit reference.')]
w=Workbook();s=w.active;s.title='Findings and Actions';s.append(['Topic','Source','Finding','Required action'])
for x in findings:s.append(x)
s=w.create_sheet('File Inventory');s.append(['File','Bytes','SHA256','Prior exact matches','Duplicate of','Pages'])
for r in inv:
 if 'sha256'in r:s.append([r['name'],r['size_bytes'],r['sha256'],len(r['prior_matches']),r['batch_duplicate_of'],r.get('pages')])
 else:s.append(['TRANSFER FAILED',r.get('error')])
s=w.create_sheet('XML Open Ends');s.append(['UID','Activity','No predecessor','No successor','Isolated'])
for t in tasks:
 no_p=t['uid']not in pred;no_s=t['uid']not in succ
 if no_p or no_s:s.append([t['uid'],t['name'],no_p,no_s,no_p and no_s])
s=w.create_sheet('XML Links');s.append(['Successor UID','Predecessor UID','XML type','Link lag'])
for l in links:s.append(list(l.values()))
s=w.create_sheet('XML Audit');s.append(['Check','Result'])
for k,v in a.items():
 if k not in ['tasks','link_records','examples','root_tags']:s.append([k,str(v)])
s.append(['no predecessor',sum(t['uid']not in pred for t in tasks)]);s.append(['no successor',sum(t['uid']not in succ for t in tasks)]);s.append(['isolated',sum(t['uid']not in pred|succ for t in tasks)])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=48
 for row in s:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
w.save(p/'BRRP_P13_Master_Plan_and_XML_Audit.xlsx')
d=Document();d.add_heading('BRRP — Master-plan, schedule XML and repeated-input review',0)
d.add_paragraph(f'P13 |25October2026 |{len(inv)} supplied file records; {sum("path"in r for r in inv)} successfully transferred; {len(unique)} distinct contents; {len(new)} distinct contents with no exact prior local match. SHA-256 comparisons identify exact duplicates; a new hash does not itself establish a newer approved revision.')
d.add_paragraph('Scope: version checking of the batch; content screening of the three distinct master-plan volumes; structural XML link/calendar/resource audit; classification of other distinct sources. Repeated sources retain prior review limitations. No claim is made of an exhaustive engineering/legal re-review, scheduler import/calculation, approved baseline or IFC/AFC release.')
for title,source,finding,action in findings:
 d.add_heading(title,1);d.add_paragraph('Source: '+source);d.add_paragraph(finding);d.add_paragraph('Required update: '+action)
d.add_heading('Next controlled issue',1);d.add_paragraph('Use one controlled source register and one governing schedule. Reconcile amendments and outstanding engineering inputs, implement approval/inspection dependencies, verify resource and vessel availability, resolve open ends, set calendars, price the full scope and run the chosen scheduler. Document critical and near-critical paths, resource overloads, monsoon exposure and cash requirements before seeking baseline acceptance. Actual procurement, fabrication, installation, commissioning and approvals remain evidence-based execution gates. Document requirements are reference content, not instructions to perform external actions.')
d.save(p/'BRRP_P13_Master_Plan_and_Schedule_Review.docx')
for name in ['BRRP_P13_Master_Plan_and_Schedule_Review.docx','BRRP_P13_Master_Plan_and_XML_Audit.xlsx']:
 q=pathlib.Path('/workspace/downloads')/name;shutil.copy2(p/name,q);q.chmod(0o644)
assert len(load_workbook(p/'BRRP_P13_Master_Plan_and_XML_Audit.xlsx').worksheets)==5
Document(p/'BRRP_P13_Master_Plan_and_Schedule_Review.docx')
print({'files':len(inv),'distinct':len(unique),'new_distinct':len(new),'findings':len(findings),'report':'saved and reopened'})
