from pathlib import Path
from datetime import date,timedelta
import json,csv,zipfile,shutil,hashlib
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt
R=Path('/workspace/project-review');O=R/'outputs/BRRP_P7_Assurance_Completion';P5=R/'outputs/BRRP_Integrated_Project_P5';G=json.loads((P5/'P5_register_data.json').read_text());D=json.loads((P5/'P5_integrated_data.json').read_text())
# Evidence-based self-review. Maturity/quality is not represented by document/activity counts.
gaps=[
('PM01','Contract-to-evidence requirements traceability','UNDERDEVELOPED','P5 source list and obligations are not a clause-by-clause verification matrix.','Contracts / Engineering','Before scope/design baseline',[
'Assign a unique requirement ID and exact source clause/revision.',
'Link each requirement to applicable bridges, calculation/drawing, procurement and test evidence.',
'Assign verification method, responsible owner and acceptance authority.',
'Close only against actual approved evidence; retain exceptions and changes.']),
('PM02','Integrated WBS / CBS / OBS / control accounts','UNDERDEVELOPED','Package IDs exist but cost, responsibility, procurement and progress are not fully mapped.','Project Controls / Cost','Before cost/schedule baseline',[
'Adopt one bridge/workstream coding dictionary across scope, schedule, cost and documentation.',
'Define control accounts with one accountable manager, budget, schedule and deliverable boundary.',
'Map procurement, material, invoice and accepted quantities to those control accounts.',
'Prevent shared marine/engineering cost duplication and cross-bridge quantity double counting.']),
('PM03','Execution-concept and constructability selection','UNDERDEVELOPED','Source concept is carried forward; project-specific alternatives and constraints are not decided.','Engineering / Marine / Construction','Before design freeze',[
'Compare practical module splits, lifting/transport strategies and fabrication routes.',
'Assess live-service clearances, existing bridge interference, access and operational restrictions.',
'Record cost, schedule, safety and interface consequences with qualified engineering input.',
'Approve the selected concept and propagate it to model, MTO, vessel needs and schedule.']),
('PM04','Engineering-model validation and checking plan','OPEN ENGINEERING','Source arithmetic and pictures do not reproduce the structural model.','Technical Authority / Checker','Before detailed design approval',[
'Issue model/input/checking registers with calculation IDs and validation responsibility.',
'Reproduce source load totals and gravity response or explain changes before new analyses.',
'Check geometry, section groups, stiffness, releases, stability, units and load combinations.',
'Document independent checking of members, joints, bearings, existing structures and temporary states.']),
('PM05','Long-lead procurement network','MATERIAL SCHEDULE GAP','P010 is one bulk-material placeholder; all25 procurement packages lack separate validated logic.','Procurement / Planning','Before schedule baseline',[
'Create separate vendor-document, approval, manufacture, inspection, dispatch and receipt milestones where relevant.',
'Obtain actual quotations and lead times for bearings, services, lifesaving equipment and specialist suppliers.',
'Connect each package to its earliest genuine design/fabrication/installation/test need date.',
'Recalculate critical paths after these links and commitments are entered; do not retain target dates by force.']),
('PM06','Schedule quality and full resource validation','UNDERDEVELOPED','Network connectivity checks pass; durations and full capacity constraints remain unverified.','Project Controls / Discipline Leads','Before baseline approval',[
'Validate every task duration against quantities, crew-hours, access and demonstrated productivity.',
'Include review cycles, vendor data, port/field moves, inspection/coating, shutdowns and operational windows.',
'Level labour, supervision, cranes, QA/NDT, coating bays, vessels, shifts, rotations and POB.',
'Audit calendars, open ends, excessive constraints, lags, safe seasonal stoppoints and milestone definitions.']),
('PM07','Activity-specific weather and operability model','CORRECTION REQUIRED','The10-day hold after offshore work is not a validated weather-risk model.','Marine / Planning / Assurance','Before offshore campaign sanction',[
'Establish operation-specific wind/wave/current/visibility limits from actual vessels and approved methods.',
'Use relevant seasonal/metocean evidence to estimate workable hours and continuous safe windows.',
'Model interruption, recovery, safe suspension, standby and demobilisation where they occur.',
'Replace the coarse hold with a documented weather/contingency basis; retain P5 only as an illustrative scenario.']),
('PM08','Quantitative schedule and cost risk assessment','NOT PERFORMED','Deterministic delay tests exist; no calibrated probabilistic risk model or confidence dates exist.','Risk / Planning / Cost','Before management forecast commitment',[
'Link risks to specific activities, costs, owners, triggers and responses.',
'Develop evidence-based duration/risk ranges and correlation assumptions.',
'Run quantitative scenarios/probabilistic assessment with documented inputs if confidence dates are required.',
'Report confidence levels and drivers only after analysis; do not invent P50/P80 dates or contingency percentages.']),
('PM09','Progress measurement and earned-value rules','UNDERDEVELOPED','Accepted quantities are specified, but discipline weighting and credit rules are not approved.','Project Controls / QA / Cost','Before first progress report',[
'Define units, denominator, milestone weights and acceptance evidence for each work type.',
'Keep physical progress, commercial billing and cash receipts separate.',
'Base project-level weighting on approved budgets/hours and control accounts; avoid arbitrary mixed-tonnage percentages.',
'Use CPI/SPI only with compatible approved planned/earned/actual baselines and consistent units.']),
('PM10','Weekly update and forecasting protocol','UNDERDEVELOPED','Reporting rhythm exists; data-date, actualisation and remaining-work rules need formal definition.','Project Controls','Before first controlled update',[
'Fix reporting data date, responsible data owners and evidence deadlines.',
'Record accepted actual starts/finishes, remaining quantities and realistic remaining durations.',
'Reconcile out-of-sequence work through authorised logic updates with preserved baseline history.',
'Issue variance narrative, critical/near-critical changes, six-week constraints and decisions with owners.']),
('PM11','Weight and centre-of-gravity control','UNDERDEVELOPED','Weight/COG is recognised but no live controlled register or change-trigger workflow exists.','Engineering / Fabrication / Marine','Before vessel commitment and each lift',[
'Assign module IDs, weight categories, uncertainty, inclusions/exclusions and COG coordinates.',
'Reconcile theoretical and measured weights including temporary steel, rigging and loose-shipped items.',
'Apply the supplied criteria reanalysis triggers:weight increase over5% or COG shift over2% of corresponding dimension.',
'Release revised lift/transport calculations and communicate changes to vessel and installation authorities.']),
('PM12','Technical interfaces and decision escalation','UNDERDEVELOPED','Interface and decision registers are described but not fully populated with closure deadlines.','Interface Manager / Project Director','Before design release and weekly',[
'Populate physical, functional, operational, organisational and document interfaces bridge by bridge.',
'Assign one interface owner, responding party, need-by date and objective closure evidence.',
'Escalate missed critical need dates and preserve decision rationale and authorised changes.',
'Separate information requests from deviations, nonconformances and commercial change requests.']),
('PM13','Supplier and subcontract performance assurance','UNDERDEVELOPED','Procurement packages exist but supplier-specific surveillance and performance controls are not developed.','Supply Chain / QA / Contracts','Before placement and during supply',[
'Qualify capability, actual capacity, competence, financial exposure and subcontract interfaces as applicable.',
'Flow down drawings/specifications, QA/HSE, documentation, schedule, change and completion obligations.',
'Agree supplier deliverable registers, inspections, expediting milestones and evidence requirements.',
'Review recovery, quality failures and counterparty risks; secure authorised remedies and alternative plans.']),
('PM14','Field change and configuration management','UNDERDEVELOPED','Change control is general; technical field-change routing and affected-document propagation need definition.','Engineering / Construction / Contracts','Before offshore execution',[
'Use distinct TQ/RFI, NCR, concession/deviation, field-change and commercial-change identifiers.',
'Require technical evaluation before any design-affecting field modification.',
'Update affected calculations, drawings, MTO, inspection, safety, schedule and cost records.',
'Maintain red-lines and superseded-document control; verbal instructions are not a design release.']),
('PM15','Critical safety-barrier verification','UNDERDEVELOPED','HSE topics are covered; operation-specific critical controls and verification evidence are not established.','HSE / Operation Owners','Before each high-risk operation',[
'Identify credible major hazards for the actual lifting, structural modification, marine and live-service tasks.',
'Assign each critical preventive/mitigating control an owner, standard and verification record.',
'Check controls at the workfront; unresolved critical barriers hold the operation.',
'Include emergency, rescue, isolation, communication and simultaneous-operation responsibilities.']),
('PM16','SIMOPS, shutdown and access integration','UNDERDEVELOPED','Access/permits are gates but operational windows are not fully scheduled and agreed.','ONGC Interface / Offshore / Marine','Before baseline and mobilisation',[
'Build platform-specific access, hot-work, isolation, shutdown and simultaneous-operation calendars.',
'Coordinate vessels, transfers, cranes, service routes and emergency access in space and time.',
'Connect agreed windows to workpacks and schedule logic; record operational restrictions.',
'Obtain renewed permissions after weather, configuration or seasonal changes.']),
('PM17','Safe seasonal suspension and restart plan','UNDERDEVELOPED','Second-season comparator assumes safe stoppoints and7-day remobilisation without operation-specific proof.','Marine / Construction / Engineering','Before adopting two-season strategy',[
'Define structurally and operationally safe suspension states for every unfinished package.',
'Engineer temporary support, preservation, service continuity and protection through the gap.',
'Plan actual demobilisation/remobilisation, reinspection, revalidation, renewed permits and costs.',
'Recalculate restart dates and commitments; seven days is only a placeholder, not established readiness.']),
('PM18','Systems-based completion and turnover','UNDERDEVELOPED','Test categories exist; detailed system boundaries and completion certificates are not mapped.','Completions / Discipline Leads / ONGC','Before construction workpacks',[
'Assign bridge/system/subsystem IDs and define structural, electrical, piping and instrument boundaries.',
'Define agreed construction-complete, ready-for-test, ready-for-service and client-acceptance evidence.',
'Link tests, punch categories, reinstatement, isolations and certificates to each boundary.',
'For BHS require accepted new bridge/services before demolition release; do not substitute structural completion alone.']),
('PM19','Operations, maintainability and asset-data handover','UNDERDEVELOPED','Dossiers are planned but operations/maintenance deliverables and asset-data requirements need agreement.','Engineering / Completions / ONGC Operations','Before detailed design freeze and handover',[
'Review bearing replacement, inspection access, drainage/coating maintenance and service access.',
'Agree applicable vendor manuals, preservation/maintenance requirements, spares and training.',
'Prepare asset/tag lists and agreed digital/as-built/model formats and ownership.',
'Obtain operations acceptance of maintainability actions and delivered asset data as applicable.']),
('PM20','Funding, contingency and financial risk governance','OPEN COMMERCIAL','Cost register is unpriced; no funding adequacy, margin or approved contingency can be claimed.','Finance / Cost / Commercial','Before financial sanction',[
'Obtain quotations and actual charter terms, mobilisation, standby, tax, FX/import/local-content treatment.',
'Build integrated cost, commitment, cash, certification, collections and bank-margin forecasts.',
'Define risk contingency ownership, drawdown authority and change/budget thresholds.',
'Test late payments, cost escalation and second-season exposures; confirm funded response options.']),
('PM21','Communication and stakeholder engagement','UNDERDEVELOPED','Meeting rhythm exists but a stakeholder/communications responsibility plan is not explicit.','Project Director / Interface Manager','At project initiation',[
'Identify ONGC engineering/operations/SCM, vessel/yard/vendor, assurance and statutory stakeholders as applicable.',
'Define authorised channels, transmittals, meeting decisions and required approval notifications.',
'Assign escalation routes for engineering, quality, operational and commercial issues.',
'Control external commitments; minutes are not automatically contract amendments or technical approvals.']),
('PM22','Data, document and delivery assurance','UNDERDEVELOPED','Files exist, but successful recipient download/opening remains unverified.','Document Control / Project Controls','Before each formal issue',[
'Use a controlled repository, permissions, revision history, backup and agreed file naming/formats.',
'Maintain issued-file manifest, checksums, transmittal purpose and affected-document links.',
'Test readability, units, formulas, print layouts and recipient access to the actual delivery route.',
'Obtain recipient acknowledgement; local file integrity does not prove successful delivery.']),
('PM23','Integrated assurance and independent review','NOT COMPLETED','Current audit is a self-review; there is no evidenced independent project/design assurance sign-off.','Project Director / Independent Review Authority','Before baseline, marine and construction release',[
'Plan risk-based reviews of scope, engineering, schedule/cost, fabrication, marine and completion readiness.',
'Assign suitably competent reviewers with required independence and clear review scope.',
'Close review findings with accountable owners and objective evidence.',
'Record release decisions and residual risks; document count and checklist completion alone do not establish adequacy.']),
('PM24','Package acceptance, lessons and closure criteria','UNDERDEVELOPED','A-to-Z actions exist; measurable package acceptance and evidence-based final closure need agreement.','Project Director / Contracts / QA','Before management package approval',[
'Agree a deliverable acceptance matrix distinguishing prepared, checked, approved, issued, executed and accepted states.',
'Link remaining P6 actions and this supplement to named owners, dates and evidence.',
'Capture actual productivity, rework and marine downtime and feed changes into forecasts.',
'Close the management package only when it is internally consistent and the required project-specific reviews are evidenced.'])]
assert len(gaps)==24
# Filled starter requirements traceability. Status is deliberately not approval.
req=[
('REQ01','Tenbridge EPC scope and interfaces','B774 structuralSOW RevC / GCC2.1.1','All10bridges','Scope matrix + bridge workpacks','Document/scope review','Project Director / ONGC as applicable'),
('REQ02','25year design life','BRRP PartII p119','Newbridges','Design basis and calculations','Engineering check','Technical Authority'),
('REQ03','Material elastic properties','PartI Annexure2 p66','Structuralmodels','Model input record E210000MPa and material properties','Model audit','Checker'),
('REQ04','Globalwalkway250kg/m2;localgrating500kg/m2','PartI bridgeclausep42;Annexure4p70','Applicablewalkways/grating','Load maps and element checks','Calculation review','Engineering/QA'),
('REQ05','Minimum1.20m clearwalkway','PartI bridgeclausep43','Newbridgeaccess','GA/clearance details and survey','Drawing/field inspection','Engineering/ONGC'),
('REQ06','Relative movement150%;hinged thrust150%','PartI bridgeclausep43','Bridge supports','Movement envelopes, bearing/connection calculations','Calculation review','Technical Authority'),
('REQ07','Applicablefuturepiping20% allowance','PartI bridgeclausep42','Piping-carrying bridges','Load register with included allowances','Calculation reconciliation','Engineering'),
('REQ08','VerticalL/400;lateralwindL/500','PartI p43;Annexure12p86','Bridges','Serviceability calculations','Model/result review','Checker'),
('REQ09','Camber compensates dead-load deflection','PartI bridgeclausep43','Newbridges','Camber profile and setting geometry','Calculation/drawing check','Technical Authority'),
('REQ10','15secgust>50m;3secgust<50m','PartI Annexure5p73','56/84/42m studyconcepts','Wind derivation and loadcases','Calculation review','Checker'),
('REQ11','TrussVIV displacement/fatigue','PartI bridgeclausep42','Applicabletrussmembers','VIV/fatigue report','Specialistcheck','Technical Authority'),
('REQ12','Actualweight/COG change-trigger reanalysis','PartI liftingcriteriap20','Liftedmodules','Weight register and revised lift analysis','Weight/liftreview','Marine/Structural'),
('REQ13','NewBHS bridge/services acceptedbefore demolition','SOW6.3.4;B774methodologyRevC','BHS-SLQ','Operationalacceptance + demolitionrelease','Operational/engineeringreview','ONGC/TechnicalAuthority'),
('REQ14','ICG/ICP/BHS strengthening addendum','Addendum_Structure_SOW p1','Applicableplatforms','Checkedwrap/landing/jacketdetails and records','Calculation/inspection','Engineering/QA'),
('REQ15','Weld/NDT/material/coatingcontrols','Specifications2009F/6001F/2005;MTO IRT','Applicablework','ApprovedITP and inspectionrecords','Inspection/audit','QA/CA/TPI as applicable'),
('REQ16','Digitalmeasurement beforebilling','TechnicalSCC e-MBprovisions','Measured/invoicedwork','Acceptedquantity/e-MB/certification/invoicepack','Commercialrecordreview','Contracts/ONGC'),
('REQ17','12months fromNOA;conditionalmonsoon relief','GCC6.3.1.1;TechnicalSCC2','Overallproject','Contractcalendar,notices and extensionrecords','Contractreview','Contracts/ONGC'),
('REQ18','Contractorconstruction supportresources','B774methodologyRevC pp137onward','Offshoreexecution','Spread,accommodation,crane/logisticsplan','Readinessreview','Marine/Construction')]
reqrows=[[x[0],x[1],x[2],x[3],x[4],x[5],x[6],'REQUIREMENT LOCATED; IMPLEMENTATION/ACCEPTANCE OPEN',None] for x in req]
# Proposed internal progress rules, not client billing percentages or actual achieved progress.
progress=[
('Engineering','Reviewed design-input/calculation scope',15,'Accepted input/checking scope'),('Engineering','Checked calculations and discipline drawings',45,'Checker records'),('Engineering','Applicable client/interface comments closed',30,'Review/closure records'),('Engineering','Authorised issue and complete records',10,'Controlled transmittal'),
('Procurement','Technically accepted order placement',10,'Authorised PO and technical acceptance'),('Procurement','Verified manufacturing milestones',30,'Agreed milestone/inspection evidence'),('Procurement','Accepted inspection and dispatch',30,'Release/dispatch records'),('Procurement','Accepted site/yard receipt',20,'Receipt/traceability'),('Procurement','Complete supplier documentation',10,'Accepted dossier'),
('Fabrication','Traceable prepared material',10,'Accepted material/preparation records'),('Fabrication','Accepted fit-up/assembly',20,'Inspection record'),('Fabrication','Completed weld production stage',35,'Weld map/records; no defect closure implied'),('Fabrication','Required NDT/repairs and dimensions accepted',20,'QA acceptance'),('Fabrication','Coating and preservation accepted',10,'Coating record'),('Fabrication','Final package release/dossier',5,'QA release'),
('Installation','Accepted landing/placement stage',25,'Approved installation record'),('Installation','Permanent connections completed',40,'Connection records; acceptance pending as applicable'),('Installation','Required inspections/restoration accepted',20,'QA acceptance'),('Installation','Package acceptance and records',15,'Accepted package closure'),
('Commissioning','Test readiness accepted',20,'Approved prerequisites and boundary'),('Commissioning','Applicable tests executed and accepted',50,'Signed test records'),('Commissioning','Reinstatement and critical punch closure',20,'Operational/QA closure'),('Commissioning','Operational acceptance and dossier',10,'Client acceptance')]
for family in set(x[0] for x in progress):assert sum(x[2] for x in progress if x[0]==family)==100
progressrows=[list(x)+['PROPOSED INTERNAL RULE; approve before use; not A-10 billing weights','No current progress percentage assigned'] for x in progress]
# Coding and responsibility mapping to existing package hierarchy.
coding=[]
for b in D['bridge_mto']:
 for family,cbs,owner in [('Engineering','ENG','Engineering Manager'),('Materials','MAT','Procurement Manager'),('Fabrication','FAB','Fabrication Manager'),('Marine','MAR','Marine Manager'),('Offshore','OFF','Offshore Manager'),('Completions','COM','Completions Manager')]:
  coding.append([b[0],family,f'BRRP/{b[0]}/{cbs}',cbs,owner,'Budget/hours and resource profile pending','Use same code in schedule/cost/PO/progress/dossier; shared costs allocated once'])
# Procurement links identify gaps without inventing supplier lead times.
links=[]
for x in G['procurement']:
 pid=x[0];need='Applicable fabrication/workpack release'
 if pid=='P07':need='Newbridge bearing/landing installation readiness (I010/I020/I030)'
 elif pid in ['P11','P12','P13','P14','P15']:need='Applicable service installation/commissioning readiness; BHS I020 sequence where relevant'
 elif pid in ['P17','P18','P23','P24']:need='Marine readiness before loadout/lifting/transport'
 elif pid=='P21':need='Each fabrication/offshore inspection holdpoint'
 elif pid=='P25':need='Beforedemolition/backload(D020/D030)'
 links.append([pid,x[1],need,None,None,None,'VENDOR LEADTIME / DOCUMENT DATES UNKNOWN','Create separate detailed network; P010 bulk placeholder does not prove this package readiness'])
# Actual confirmed issues, separate from generic empty templates.
issues=[
('TQ01','Technical query','BHS SLQ support wording','SOWcontradiction vs report/methodology fixed intent','Engineering/ONGC','Approvedsupportphilosophy and drawing/modelupdate'),
('TQ02','Technical query','Camber rule versus source1:200','Suppliedcriteria requires dead-load compensation','TechnicalAuthority','Checkedcamberprofile and instructions'),
('TQ03','Technical query','Live-load totals and actual areas','ICW156.25vs144.69kN;otherrectangular screens differ','AnalysisLead','Reconciledloadmap and modelrevision'),
('MTO01','Quantity reconciliation','Structural1141vs1142MT','Landingstrengthening216vs217MT','Engineering/QS','ControlledMTOreconciliation'),
('MTO02','Quantity reconciliation','Handrail1166vs1168m','Source rows versus printedtotal','Engineering/QS','Drawing-basedtake-off'),
('MTO03','Quantity reconciliation','Grating718vs719m2','Source rows versus printedtotal','Engineering/QS','Drawing-basedtake-off'),
('SCH01','Planning correction','Longlead packages not individually validated','P010generic25daymaterialsworkflow','Procurement/Planning','Vendor-backednetworkupdate'),
('SCH02','Planning correction','Weatherreserve model too coarse','10dayholdafterproductivework','Marine/Planning','Operation-specificweather/contingencymodel'),
('DEL01','Delivery issue','Recipientdownloadsnotconfirmed','User repeatedlyreportedfileaccessfailure','DocumentControl','Recipientopensactualissuedfiles')]
# Commissioning/turnover boundaries to be agreed, not invented client certificate terminology.
systems=[
('SYS-STR','Eachbridge permanentstructure','Connections,grating,handrails,coating,dimensions','Construction/QA','Acceptedinspections and punchstatus','Required for bridge acceptance'),
('SYS-BRG','Newbridge bearings/guides','Settings,movementclearances,restraints,protection','Engineering/QA','Acceptedinstallation/setting records','Required before operational acceptance'),
('SYS-ELE','Applicableelectricalcircuits','Cables,trays,lighting,earthing,isolations','Electrical/ONGC','Approvedtests and energisation/reinstatement','Boundaries/values to agree'),
('SYS-PIP','ApplicableBHS pipe systems','Lines,tie-ins,supports,tests,isolation/reinstatement','Piping/ONGC','Line-specifictests and operatingacceptance','Required before relevantoldservice removal'),
('SYS-INS','Applicableinstrument/controlloops','Loops,functions,isolations and restoration','Instrumentation/ONGC','Loop/functiontestrecords','Applicable scope to confirm'),
('SYS-LSA','Applicablelifesaving/accessinterfaces','Muster,access,lifeboatinterfaces,safetystudyactions','HSE/ONGC','Approvedcommissioning/operationalrecords','Scope/vendorrequirements to confirm'),
('SYS-DMO','BHS existingbridge removal','Isolation,stability,demolition/backload','TechnicalAuthority/ONGC','Newbridge/servicesaccepted + demolitionrelease','No removal on structuralplacementalone')]
# Clear maturity statement: self-review, no performance ranking or certifications.
status=[['Benchmark claim','NOT SUBSTANTIATED','No audited comparison withSaipem/McDermott; companyinternalprocedures and outcomes not verified.'],['Review independence','SELF-REVIEW','This supplement is not independentengineering or projectassurance sign-off.'],['Management documents','PRELIMINARY / DEVELOPING','P5/P6 provide scope,registers and workstreamplans; some procedures require project-specificpopulation/review.'],['Engineering','INCOMPLETE','P4source-basedsheets/arithmetic are not completedglobal/member/connection/bearing/fatigue/liftingdesign.'],['Schedule','ILLUSTRATIVE CONDITIONAL SCENARIO','Validconnectivity does not establish scopecoverage,resourcefeasibility or weatherconfidence.'],['Quantities','TENDER / UNRECONCILED','Three arithmeticdiscrepancies remainopen; nofinalpurchaseMTO.'],['Cost/cash','UNPRICED','Noapprovedbudget,margin,fundingadequacy or costconfidence.'],['AFC','NOT APPROVED','Register/forms are not actualapproveddrawings.'],['Execution/acceptance','NOT EVIDENCED HERE','No suppliercommitment,physicalcompletionorclientacceptance inferred.'],['Delivery','NOT CONFIRMED','Workspacefileintegrity does notprove recipientdownload/opening.']]
reg=[]
for id,title,kind,finding,owner,gate,actions in gaps:
 for i,action in enumerate(actions,1):reg.append([f'{id}.{i:02}',id,title,kind,finding,action,owner,gate,'OPEN / PROJECT EVIDENCE REQUIRED'])
assert len(reg)==96
# Shared fields for actual action tracking; not invented claims or performed approvals.
workflows=[['RFI/TQ','Clarify existing information/requirement','Engineeringresponse;recordsource;assessaffectedwork','No automaticdesignchangeauthorisation'],['NCR','Actualwork/materialdoesnotmeetrequirement','Contain;evaluate;approvedrepair/rework/disposition;verifyclosure','Concessionrequiresauthorisedtechnical/clientacceptancewhereapplicable'],['Deviation/concession','Proposeddeparturefromrequirement','Technical/commercialevaluation;approvalbeforeuse','Do not treatassilentscopechange'],['Fieldchange','Proposedchangeatsite','Technicalreviewandaffectedcalculation/drawing/MTO/safetyreview','Workhelduntilrequiredrelease'],['Commercialchange','Potentialcontractscope/time/costchange','Timelynotices;instruction;substantiation;authorisedsettlement','Noentitlementassumedfromtechnicalqueryalone']]
# Workbook.
w=openpyxl.Workbook();w.remove(w.active)
def tab(name,head,rows,widths=None):
 s=w.create_sheet(name);s.append(head)
 for r in rows:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
 for i in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(i)].width=widths[i-1] if widths and len(widths)>=i else 35
 for row in s:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 return s
tab('Actual Maturity',['Area','Status','Meaning'],status,[32,55,125]);tab('Review Findings',['ID','Control','Classification','Evidence-basedgap','Owner','Needby'],[list(x[:6]) for x in gaps],[15,60,34,110,40,45]);tab('96 Closure Actions',['ID','Finding','Control','Class','Observedgap','Requiredaction','Owner','Needby','Status'],reg,[18,14,60,28,100,115,40,45,45]);tab('Requirements Traceability',['RequirementID','Requirement','Source','Applicability','Deliverable/evidence','Verification','Acceptanceauthority','Status','Evidence/reference'],reqrows,[20,65,65,35,80,40,45,60,40]);tab('Control Account Map',['Bridge','Workstream','Proposedcontrolaccount','CBScategory','Accountablemanager','Budgetbasis','Application'],coding,[26,30,50,20,40,55,105]);tab('Progress Rules Proposal',['Discipline','Milestone','Proposedweightpercent','Requiredevidence','Approvalstatus','Actualprogress'],progressrows,[25,65,28,65,95,55]);tab('Procurement Network Gaps',['PackageID','Package','Requiredneedpoint','Vendorleadtime','Vendordocapprovaldate','Deliverydate','Status','Networkaction'],links,[18,55,105,25,30,30,60,110]);tab('Technical and Planning Issues',['ID','Type','Issue','Evidence','Owner','Closure'],issues,[18,32,70,90,40,80]);tab('Turnover Boundaries',['SystemID','Boundary','Includedscope','Owner','Acceptanceevidence','Dependency'],systems,[20,50,90,40,100,80]);tab('Change Workflows',['Type','Trigger','Process','Control'],workflows,[25,70,110,110])
w.save(O/'BRRP_P7_Package_Assurance_and_Closure.xlsx')
# Word supplement with explicit procedures and proposed rules.
d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10);d.add_heading('BRRP — Project Management Package Self-Review and Completion Supplement',0);d.add_paragraph('P7 |25October2026 | Forreview | No claim of superiority, independentassurance or constructionapproval')
d.add_paragraph('Conclusion: the package cannot be substantiated as better than leading offshore EPC companies. It remains a preliminary planning/engineering-development package. Document counts, correctly calculated example arithmetic and a connected schedule do not demonstrate a complete, executable or approved project. This review identifies24additional or insufficientlydeveloped controls and96closure actions. It supplements, rather than duplicates or cancels, the104P6actions.')
def table(head,rows,size=8):
 t=d.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,head):c.text=str(v)
 for r in rows:
  for c,v in zip(t.add_row().cells,r):
   c.text='' if v is None else str(v)
   for p in c.paragraphs:
    for run in p.runs:run.font.size=Pt(size)
 return t
d.add_heading('1. Actual maturity and correction to any overstatement',1)
table(['Area','Status','Meaning'],status)
d.add_paragraph('A101-itemengineeringregister does notmean101completeddesigns/drawings. The nineP4development sheets comprise source-informedarrangements,functionalsupportinterfaces and detailing schedules; they are not ninefullyengineereddrawings. P5/P6remainusefulpreparation, but the managementpackage cannot be calledbaseline-ready while engineering, costs, resourcecommitments, procurementlogic, weather and reviewevidence are unresolved.')
d.add_heading('2. Evidence-based schedule findings',1)
d.add_paragraph('The existing205activityscenario has one start and one terminal milestone; dependency/calendar/exclusive-namedresourcechecks passed. This validates those internal properties only. The25procurementpackages are not individuallyrepresented with vendor-backedlead times; P010contains onegeneric25daybulk-materialworkflow. Threecrews, yard/QA/coating/POB capacity and mostduration/productivityinputs remainassumptions. The10dayweatherholdafterproductivework is not an operation-specificweatherforecast. The15June2027date must remainillustrative until these gaps are closed; finalprojectcriticalpaths may change.')
d.add_paragraph('The earlier3daymaterialdelaytest is a sensitivity of the coarsecalendar/reserveassumption, not proof that realdemobilisation mustwaituntilOctober. Finalplanningmustuseengineeredoperatinglimits, actualweather/workability, safesuspension and demobstrategies. Contractualrelief and costsremainseparateassessment.')
d.add_heading('3. Additional or underdeveloped controls',1)
for id,title,kind,finding,owner,gate,actions in gaps:
 d.add_heading(id+' — '+title,2);d.add_paragraph('Classification: '+kind+'. Finding: '+finding+' Owner: '+owner+'. Need-by: '+gate+'.')
 for i,t in enumerate(actions,1):d.add_paragraph(f'{id}.{i:02} {t}')
 d.add_paragraph('Closure: project-specificprocedure/register populated, reviewed by designatedauthority, linked to objective evidence and reflected in affected baseline/deliverables. Preparationalone is not implementationapproval.')
d.add_heading('4. Requirement-to-evidence traceability: populated starter matrix',1)
table(['ID','Requirement/source','Deliverable/evidence','Verification/authority'],[[x[0],x[1]+'; '+x[2],x[4],x[5]+'; '+x[6]] for x in reqrows])
d.add_heading('5. Integrated coding and control-account procedure',1)
d.add_paragraph('Proposedcode:BRRP/<bridge>/<workstream>. WorkstreamsENG,MAT,FAB,MAR,OFF,COMmap engineering,materials,fabrication,marine,offshore and completions. Eachcontrolaccount needs oneaccountablemanager, approvedscope, deliverables,budget/hours, schedule and progressrules. POitems, materialissues, timesheets, acceptedquantities,invoices and turnoverrecords mustcarrycompatiblecodes. Sharedcosts mustuse a controlledallocationonce; they mustnot be countedbothcentrally andagainbybridge.')
d.add_heading('6. Proposed progress-measurement procedure',1)
d.add_paragraph('The following are proposed internal physical-progress rules, not contractbillingmilestones. Approve them beforeuse and align with acceptedwork and budget/hours. Do not assignactualprogress without evidence. Disciplineweights total100%; projectlevelweights require an approvedbudget/hourbasis. Do notaveragebridgeMT, drawingcounts, services and cash as if they were oneunit.')
table(['Discipline','Milestone','Proposedweight%','Evidence'],[[x[0],x[1],x[2],x[3]] for x in progressrows])
d.add_paragraph('Weightedprogress=SUM(approvedmilestoneweight x evidencedfraction). Stagefractionsuseconsistentunits andacceptance; preservepriorapprovedweights. EV=approvedbudgetedvalueofperformedwork; AC=actualcostforthesamescope/period/currency; PV=timephasedapprovedplan. CPI=EV/AC and SPI=EV/PV only when denominators/baselines are valid. Currentcost/earnedvalue inputs are notestablished; no projectCPI/SPI isissued.')
d.add_heading('7. Controlled weekly-update and decision cycle',1)
for t in ['Fixa singledata date and authoriseddata owners. Collect actualstarts/finishes, acceptedquantities, inspectionstatus, remainingwork and actualresource/costrecords.','Validateactuals and reconcile out-of-sequencework; preservebaseline and approvedchangehistory. Do notforceforecastdatestomatchatarget.','Recalculatecritical/near-criticalpaths, forecastcost/cash, weatherexposure and procurement/resourceconstraints.','Escalate when a critical prerequisite misses itsrequired-bydate, an operationalgate isHOLD, a supplierwindow isuncommitted or a payment/fundingrequirement isuncovered. Internalresponsetimes mustbe approved, not misrepresented as contractlimits.','Issue weeklyreport with decisions, owner, need-bydate, evidence and consequence. Update the six-weeklook-ahead usingonlyexecutableworkfronts.']:
 d.add_paragraph(t,style='List Bullet')
d.add_heading('8. Technical-query, nonconformance and change workflows',1)
table(['Recordtype','Trigger','Process','Releaseboundary'],workflows)
table(['IssueID','Issue','Owner','Requiredclosure'],[[x[0],x[2],x[4],x[5]] for x in issues])
d.add_heading('9. Procurement and resource integration',1)
d.add_paragraph('The appendedworkbookmaps all25packages to needpoints while leaving unknownlead times/datesblank. Obtainvendor/subcontractdeliveryandapprovalmilestones; addnecessaryengineering/vendor-data/inspection/logisticsactivities; linkreceipttoactualworkfrontreadiness. Validateallsharedresources, includinginspection/coating, supportvessels, transfers and POB, notonlyfabricationlines/threecrews. Noorder,charterorcontactwithothers isperformedbythissupplement.')
d.add_heading('10. Completion/turnover and operations acceptance',1)
table(['BoundaryID','Scope','Evidence','Dependency'],[[x[0],x[2],x[4],x[5]] for x in systems])
d.add_paragraph('Agree clientterminology and exactcertificatecriteria for constructioncomplete, readyfortesting, readyforservice and contractualacceptance. These labels are planningcategories, notassertedcontractterms. Testvalues, punchcategories, reinstatement/isolations, manuals/spares/training and assetdata requirements mustbe project-specificandapproved. BHSdemolitionremainsdependent on acceptednewbridge/services and engineeredremovalrelease.')
d.add_heading('11. Gate for declaring the management package complete',1)
for t in ['Every applicablecontract/scope/designrequirement maps to a deliverable, verificationmethod, owner and acceptanceauthority.','Scope,MTO,engineering,procurement,schedule,resources,cost/cash,methods,quality/HSE and turnover are internallyconsistent.','Project-specificinputs and uncertainties are resolved or explicitlyaccepted through authorised risk/exception decisions; missingtechnicaladequacy is nothiddenby a managementexception.','Reviewfindings have actualclosureevidence; suitablyindependentreviews are performedwhere required.','Issuedfiles are readable and successfullyreceived/opened through an agreed deliveryroute; localchecksums alone are insufficient.','Actualengineeringapproval,AFCrelease,suppliercommitments,physicalwork and clientacceptance are statedonlywhenevidenced.']:
 d.add_paragraph(t,style='List Bullet')
d.add_paragraph('Nextstep:populate project-specificinputs and closepriorityschedule/engineering/commercialgaps; completecompetentreviews. Thisself-reviewstrengthensthepackagebutdoesnotcertifyprojectcompleteness or establishcompanyperformance superiority.')
d.save(O/'BRRP_P7_Management_Package_Review_and_Remaining_Details.docx')
with open(O/'BRRP_P7_96_Closure_Actions.csv','w',newline='') as f:
 z=csv.writer(f);z.writerow(['ActionID','FindingID','Control','Class','Finding','Action','Owner','Needby','Status']);z.writerows(reg)
(O/'P7_review_data.json').write_text(json.dumps({'findings':gaps,'actions':reg,'requirements':reqrows,'progress':progressrows,'control_accounts':coding,'procurement_links':links,'issues':issues,'turnover':systems,'maturity':status},indent=2))
print('Prepared24findingcontrols/96actions;18requirementtraceabilityrows;60controlaccountmappings;25procurementlinks;proposedprogressrules;9actualtechnical/planning/deliveryissues;7turnoverboundaries. Self-reviewonly.')
