from pathlib import Path
import math,json,csv
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt,Inches
from zipfile import ZipFile
O=Path('/workspace/project-review/outputs/design_development')
rows=[['ICW-ICG LD','B774-000-83-43-BPR-0004 Rev 0',56,54.5,1.2,105,87,.557,.509,85,'133-135;152-153'],['ICP-ICG LD','B774-000-83-43-BPR-0003 Rev 0',84,80.5,1.2,177,123,.610,.625,143,'196-199;207-209'],['BHS-SLQ LD','B774-000-83-43-BPR-0005 Rev A',42,40,2.324,32,76,.916,.859,32,'97-100;122-123']]
wb=openpyxl.Workbook();wb.remove(wb.active)
def sheet(name,headers,data,widths=None):
 s=wb.create_sheet(name);s.append(headers)
 for r in data:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.fill=PatternFill('solid',fgColor='17365D');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True)
 for r in s.iter_rows(min_row=2):
  for c in r:c.alignment=Alignment(wrap_text=True,vertical='top')
 for i in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(i)].width=widths[i-1] if widths and len(widths)>=i else 22
 return s
basis=[['Status','PRELIMINARY DESIGN DEVELOPMENT / INDEPENDENT ARITHMETIC AUDIT ONLY. Not a completed design, checked FE model or AFC release.'],['Sources','SOW & Bridge Report Combined.pdf, 506 pages; structural SOW Rev C; reports BPR-0003/0004/0005; structural scope addendum; Specification Combined.pdf.'],['Scope','Three replacement bridges and associated supports / landings / platform-jacket interfaces. BHS includes service supports, operational transfer and demolition engineering.'],['Geometry','Reported SACS lengths 56/84/42 m are preliminary model lengths, NOT verified fabrication lengths or bearing-centre coordinates.'],['Loads','Load extents, source model and combinations require reconciliation; do not size members from the arithmetic review.'],['Limits','L/400 vertical and L/500 lateral are source-study criteria used for arithmetic reproduction, not independently approved project acceptance criteria.'],['Sections','Tubular dimensions in Section Geometry are transcribed only from ICP report VIV table. They are NOT mapped to chords/braces or selected for construction.'],['Support','SOW fixed/sliding terminology must be implemented using verified global/local directions, support coordinates, movements and actual restraints. Fixed is not automatically a moment-fixed member end.'],['Model status','Native analysis model, member groups, load distribution, stiffness assumptions, connectivity, combinations and full result envelopes have not been independently reproduced.'],['Release','Final fabrication, bearing, connection, lifting and temporary-work details require checked project-specific calculations and applicable ONGC approval.'],['Spreadsheet','Excel formula fields are recalculated on opening. Audit Snapshot contains independently computed static values for immediate reading.']]
sheet('Basis',['Topic','Statement'],basis,[28,140])
s=sheet('Source Inputs',['Bridge','Source report','Model length m','Approx gap m','Walkway loading width m','Reported vertical mm','Reported lateral mm','Highest listed member UC','Highest listed joint UC','Reported DL deflection mm','Combined PDF pages'],rows,[22,42,20,20,25,23,23,25,25,26,28])
calc=[];snap=[]
for i,r in enumerate(rows,2):
 bridge,doc,L,gap,b,dv,dh,uc,juc,dd,pages=r
 calc.append([bridge,f"='Source Inputs'!C{i}",f'=B{i}*1000/400',f"='Source Inputs'!F{i}",f'=D{i}/C{i}',f'=B{i}*1000/500',f"='Source Inputs'!G{i}",f'=G{i}/F{i}',f'=B{i}*1000/200',f"='Source Inputs'!J{i}",f'=I{i}/J{i}','SOURCE LIMIT ARITHMETIC ONLY; not independent structural adequacy'])
 snap.append([bridge,L*1000/400,dv,dv/(L*1000/400),L*1000/500,dh,dh/(L*1000/500),L*1000/200,dd,(L*1000/200)/dd])
sheet('Serviceability Arithmetic',['Bridge','Model L m','Source L/400 mm','Reported vertical mm','Vertical ratio','Source L/500 mm','Reported lateral mm','Lateral ratio','L/200 arithmetic mm','Reported DL deflection mm','L/200 divided by DL deflection','Interpretation'],calc)
sheet('Audit Snapshot',['Bridge','Source vertical limit mm','Reported vertical mm','Ratio vertical','Source lateral limit mm','Reported lateral mm','Ratio lateral','L/200 arithmetic mm','Reported DL deflection mm','Camber arithmetic / DL'],snap)
loads=[]
for r,table,output in zip(rows,[156.25,288,268.42],[144.69,288,268.42]):
 L=r[2];b=r[4];q=250*9.80665/1000;w=L*b*q
 loads.append([r[0],L,b,250,q,w,table,output,table-w,'Screen uses full rectangular L x width; source may include/exclude other areas. Obtain actual load map. No discrepancy correction assumed.'])
sheet('Live Load Reconciliation',['Bridge','Study L m','Loading width m','Source loading kg/m2','Converted kN/m2','Simple rectangular load kN','Report table LL kN','Output LL magnitude kN','Table minus rectangle kN','Required interpretation'],loads)
sections=[]
for od,t,length in [(406,15.9,4),(406,12.7,4),(168,9.5,4.98),(114,7.9,4.71)]:
 inside=od-2*t;A=math.pi/4*(od**2-inside**2);I=math.pi/64*(od**4-inside**4);Z=I/(od/2);rad=(I/A)**.5;mass=A*.00785
 sections.append([od,t,inside,A,I,Z,rad,mass,length,'ICP VIV table, combined PDF p.208; member role / grade / corrosion / effective length not adopted'])
sheet('Section Geometry',['Source OD mm','Source wall mm','Calculated ID mm','Area mm2','I mm4','Elastic Z mm3','Radius gyration mm','Nominal mass kg/m','Source VIV length m','Status'],sections)
queries=[
['Q01','Geometry','All','Model length versus final span','Survey coordinates, datums, skew, landing offsets and actual bearing centres','HOLD','Engineering'],
['Q02','Analysis','All','Source model unavailable for reproduction','Native SACS model, connectivity, groups, loads, combinations, stiffness and envelopes','HOLD','Analysis lead'],
['Q03','Loads','All','Rectangular LL screen does not reproduce all totals','Actual loaded areas, extensions, units, distribution and load-generation records','HOLD','Analysis lead'],
['Q04','Loads','ICW','Report LL156.25kN vs output144.69kN','Identify controlling model revision and corresponding load map','HOLD','Analysis lead'],
['Q05','Camber','All','Report instruction 1:200 differs from reported DL deflection','Define camber meaning/profile, reference line, support levels and tolerances','HOLD','Technical Authority'],
['Q06','Support','BHS-SLQ','Report says SLQ fixed; SOW also contains sliding wording','Resolve support philosophy and exact freedoms / restraints','HOLD','ONGC / Engineering'],
['Q07','Wind','All','15-second gust in IC reports vs 3-second gust in BHS','Approved metocean/design basis and compatible averaging/conversion method','HOLD','Analysis lead'],
['Q08','Existing structure','All','Bridge reports refer to separate support/deck/jacket adequacy','Obtain checks/models and integrate reactions with addendum strengthening','HOLD','Platform engineer'],
['Q09','Reactions','All','Summary does not identify complete concurrent signed envelopes','Provide units, axes, governing combinations, minima, maxima, uplift and moments','HOLD','Analysis lead'],
['Q10','Members','All','Pictures/VIV list do not establish member schedule','Obtain complete section-group mapping, effective lengths and material specifications','HOLD','Analysis lead'],
['Q11','Dynamic/fatigue','All','Source summary does not establish full project coverage','Verify vibration, wind response, fatigue and movement criteria with full calculations','HOLD','Specialist engineer'],
['Q12','Temporary conditions','All','Operating study is not lifting / transport approval','Check actual modules, transport supports, accelerations, lifting and installation states','HOLD','Marine / Structural'],
['Q13','References','All','83-41 / 83-43 and landing/demolition reference conflicts','Issue corrected controlling drawing register and revisions','HOLD','Document control / ONGC'],
['Q14','Durability','All','50ksi yield is not a complete procurement grade','Approve grade, toughness, thickness tolerances, corrosion basis, welding and coatings','HOLD','Materials / QA']]
sheet('Design Queries',['ID','Discipline','Bridge','Issue','Required closure','Status','Owner'],queries,[10,24,20,55,90,16,25])
checks=[['Global model','Permanent framing','Coordinates, member connectivity, releases, eccentricity, torsion and lateral stability','Checked model and independent review'],['Members','Strength and buckling','Axial, bending, shear, interaction, local slenderness and effective lengths','Member group schedule and full utilisation envelopes'],['Connections','Joints and end connections','Tubular joint checks, local stresses, bolts, welds, gussets and fatigue details','Joint calculations and fabrication detail drawings'],['Serviceability','Deflection and vibration','Agreed load cases, vertical/lateral limits, walkway vibration and service movement','Serviceability report and camber profile'],['Bearings','Restraints and movements','Thermal and relative platform motion, friction, guides, travel and uplift','Bearing data sheets and reactions'],['Existing steel','Landings / decks / jackets','Actual corrosion and load paths plus new reaction envelopes','Integrated existing-structure adequacy report'],['Marine','Temporary conditions','Loadout, transport, seafastening, lifting, landing and temporary support','Checked operation-specific engineering'],['Fabrication','Detailed engineering','Approved geometry and sections, connections, tolerances, inspection and traceability','AFC drawings and manufacturing dossier']]
sheet('Calculation and Detail Plan',['Workstream','Scope','Required analysis','Release deliverable'],checks,[25,30,100,65])
wb.calculation=openpyxl.workbook.properties.CalcProperties(calcId=191029,fullCalcOnLoad=True)
wb.save(O/'BRRP_Preliminary_Engineering_Review_RevP2.xlsx')
# Engineering review Word report with independently evaluated results.
d=Document();d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
d.add_heading('BRRP — Preliminary Structural Engineering Review',0)
d.add_paragraph('Rev P2 | 25 October 2026 | Source-based design development and arithmetic audit')
d.add_paragraph('This package reviews the uploaded replacement-bridge studies; it is not a newly completed or certified structural design. No structural analysis model has been independently reproduced. No fabrication, bearing, connection, lifting or demolition detail is approved by this review.')
d.add_heading('1. Adopted source basis',1)
for x,y in basis[:10]:d.add_paragraph(x+': '+y)
d.add_heading('2. Reported structural results and arithmetic checks',1)
t=d.add_table(rows=1,cols=7);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Bridge','Study L m','Listed member UC','Listed joint UC','Vertical mm / source limit','Lateral mm / source limit','Source PDF pages']):c.text=v
for r,a in zip(rows,snap):
 vals=[r[0],str(r[2]),str(r[7]),str(r[8]),f'{r[5]} / {a[1]:.0f} ({a[3]:.1%})',f'{r[6]} / {a[4]:.0f} ({a[6]:.1%})',r[10]]
 for c,v in zip(t.add_row().cells,vals):c.text=v
d.add_paragraph('These values reproduce source-study summaries and their stated L/400 and L/500 arithmetic. They are not new analysis results or acceptance of current adequacy. Listing a UC below 1.0 does not verify missing load cases, platform capacity, fatigue, temporary states or changed geometry.')
d.add_heading('3. Camber clarification',1)
for r,a in zip(rows,snap):d.add_paragraph(f'{r[0]}: literal L/200 gives {a[7]:.0f} mm; reported dead-load deflection is {a[8]:.0f} mm. Obtain the intended fabrication camber profile and reference geometry. No camber is selected here.')
d.add_heading('4. Live-load consistency screen',1)
d.add_paragraph('Using 250 kg/m2 as mass loading gives 2.4516625 kN/m2. A simple full-length rectangular walkway calculation is a screening comparison only. Real loaded areas, landing extensions and source model generation can differ.')
for r in loads:d.add_paragraph(f'{r[0]}: L x b x q = {r[5]:.2f} kN; report table {r[6]:.2f} kN; reported output {r[7]:.2f} kN. Reconcile source areas and model revision before sizing.')
d.add_heading('5. Source section geometry — not a member selection',1)
d.add_paragraph('ICP report VIV table lists 406 x 15.9, 406 x 12.7, 168 x 9.5 and 114 x 7.9 mm tubular sections. Their member functions cannot be assigned from that table. Workbook geometry uses A = pi(D^2-d^2)/4, I = pi(D^4-d^4)/64, Z = 2I/D, r = sqrt(I/A), nominal mass = A x 0.00785 kg/m. No corrosion deduction, tolerance, effective length, resistance or member selection is implied.')
d.add_heading('6. Required technical-query closures',1)
for q in queries:d.add_paragraph(f'{q[0]} | {q[2]} | {q[3]}. Required: {q[4]}. Owner: {q[6]}.')
d.add_heading('7. Detailed-engineering work programme',1)
for r in checks:d.add_paragraph(r[0]+' — '+r[2]+'. Deliverable: '+r[3]+'.')
d.add_paragraph('For every replacement bridge, develop GA, framing, secondary steel, bearings, end connections, landing interfaces, jacket strengthening, access fittings, service supports, temporary/lifting details, transport supports and MTO/inspection drawings. Add BHS service-transfer and demolition details. Each drawing must identify the supporting approved calculation, survey revision and design status.')
d.add_heading('8. Release position',1)
d.add_paragraph('The source studies provide a credible starting concept: 56m, 84m and 42m braced space-truss models. Their study-level conclusions are not independently adopted. Proceed with input verification, native-model review and design development. AFC release remains subject to competent project-specific checking and applicable ONGC approval.')
d.save(O/'BRRP_Preliminary_Engineering_Review_RevP2.docx')
# static machine-readable audit to avoid reliance on cached Excel formula values
with open(O/'BRRP_Engineering_Audit_Static.csv','w',newline='') as f:
 w=csv.writer(f);w.writerow(['Bridge','Vertical limit mm','Reported vertical mm','Ratio vertical','Lateral limit mm','Reported lateral mm','Ratio lateral','Literal L/200 mm','Reported DL mm','Camber ratio']);w.writerows(snap)
# Validate independently recomputed figures and office container integrity.
assert all(a[3]<1 and a[6]<1 for a in snap)
assert abs(snap[1][3]-177/210)<1e-12
assert abs(snap[2][6]-76/84)<1e-12
for p in O.iterdir():
 if p.suffix in ['.xlsx','.docx']:
  with ZipFile(p) as z:assert z.testzip() is None
print('SOURCE-BASED AUDIT:')
for x in snap:print(x)
print('LIVE LOAD SCREENS:')
for x in loads:print(x[:9])
print('Created workbook (8 sheets), Word review, static CSV; validated arithmetic and file integrity. No independent FE analysis or AFC design claimed.')
