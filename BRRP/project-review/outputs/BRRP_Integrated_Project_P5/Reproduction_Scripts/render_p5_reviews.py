from pathlib import Path
from datetime import date,timedelta
import json,math,shutil,zipfile,hashlib
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3,A4,landscape
from reportlab.lib.colors import HexColor,black,white
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,PageBreak
from reportlab.lib.styles import getSampleStyleSheet,ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from docx import Document
from docx.shared import Pt
import fitz
import runpy
clean=runpy.run_path('/workspace/project-review/beautify_p5.py')['clean']
R=Path('/workspace/project-review');O=R/'outputs/BRRP_Integrated_Project_P5';D=json.loads((O/'P5_integrated_data.json').read_text());M=json.loads((O/'P5_publication_meta.json').read_text());G=json.loads((O/'P5_register_data.json').read_text());S=D['schedule'];by={x['id']:x for x in S};START=date.fromisoformat(D['start'])
# Overall Gantt (working intervals, not phase completion boundary dates).
W,H=landscape(A3);p=O/'BRRP_P5_Overall_Project_Schedule.pdf';c=canvas.Canvas(str(p),pagesize=(W,H));c.setTitle('BRRP P5 Overall Project Schedule - Conditional Scenario')
def tx(x,y,t,size=9,col=black):c.setFont('Helvetica',size);c.setFillColor(col);c.drawString(x,y,clean(str(t)))
tx(35,H-40,'SHIVANSH / ONGC BRRP — OVERALL PROJECT SCHEDULE P5',16,HexColor('#17365D'))
tx(35,H-62,'PROPOSED SCENARIO | assumed NOA1Nov2026 |12months contractual basis | durations/resources unverified',10)
tx(35,H-80,'Red: phase with critical child work. Blue: other work. Shaded:16May–15Oct offshore closure. Bars show last workdays.',9)
left=490;right=W-35;scale=(right-left)/400;top=H-125;bottom=110
# Background offshore closure band
closure0=(date(2027,5,16)-START).days;closure1=(date(2027,10,16)-START).days
c.setFillColor(HexColor('#FFF2CC'));c.rect(left+closure0*scale,bottom,(closure1-closure0)*scale,top-bottom,fill=1,stroke=0)
for year,month in [(2026,11),(2026,12)]+[(2027,m) for m in range(1,13)]:
 dt=date(year,month,1);x=left+(dt-START).days*scale
 if x<=right:
  c.setStrokeColor(HexColor('#D0D0D0'));c.line(x,bottom,x,top);tx(x+2,top+10,dt.strftime('%b %y'),8)
phases=M['phase_gantt'];rowh=min(14,(top-bottom-22)/len(phases));tx(35,top+10,'PHASE / PACKAGE',9);tx(330,top+10,'START — LAST WORKDAY',9)
for k,(pid,s,f,crit,pkg) in enumerate(phases):
 y=top-15-k*rowh;name=by[pid]['name'].replace(' - phase complete','') if pid in by else pid
 tx(35,y,pid+' | '+pkg+' | '+name[:33],8);tx(330,y,str(START+timedelta(days=s))+' — '+str(START+timedelta(days=f-1)),8)
 c.setFillColor(HexColor('#C0504D') if crit else HexColor('#4F81BD'));c.rect(left+s*scale,y-2,max((f-s)*scale,2),7,fill=1,stroke=0)
for label,dt,color in [('Internal handover',date(2027,6,15),'087E8B'),('Unextended deadline',date(2027,10,31),'A00000')]:
 x=left+(dt-START).days*scale;c.setStrokeColor(HexColor('#'+color));c.setDash(3,3);c.line(x,bottom,x,top);c.setDash();tx(x-25,92,label,8,HexColor('#'+color));tx(x-25,80,str(dt),8)
tx(35,58,'Critical route: materials → BHS fabrication/installation/services → BHS demolition → ICP installation → offshore closure → dossiers/acceptance.',9)
tx(35,43,'Calendar sensitivity is significant. Validate actual NOA, workability, safe seasonal suspension, procurement lead times and all resource commitments.',9)
c.showPage();c.save()
# Detailed CPM review PDF with source/quantity/schedule evidence.
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='SmallP',fontName='Helvetica',fontSize=8,leading=10,spaceAfter=4))
styles.add(ParagraphStyle(name='TinyP',fontName='Helvetica',fontSize=7,leading=9))
story=[]
def P(t,style='BodyText'):story.append(Paragraph(clean(t).replace('&','&amp;'),styles[style]));story.append(Spacer(1,5))
def T(head,rows,widths):
 content=[[Paragraph(clean(str(v)).replace('&','&amp;'),styles['TinyP']) for v in head]]+[[Paragraph(clean(str(v)).replace('&','&amp;'),styles['TinyP']) for v in row] for row in rows]
 t=Table(content,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),HexColor('#D9E2F3')),('GRID',(0,0),(-1,-1),.3,HexColor('#888888')),('VALIGN',(0,0),(-1,-1),'TOP'),('LEFTPADDING',(0,0),(-1,-1),4),('RIGHTPADDING',(0,0),(-1,-1),4),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]));story.append(t);story.append(Spacer(1,10))
P('BRRP — Integrated Project and Critical-Path Review P5','Title')
P('25 October2026 | Preliminary / for review | Construction release deferred','SmallP')
P('1. What is prepared','Heading1')
P('The package covers10bridge scopes,100work packages,101engineering/AFC-review deliverables,25procurement packages,205schedule activities/milestones, quality/HSE gates, commissioning, cost-input controls and handover. Engineering source reviews and P4development sheets are carried forward. Procurement, fabrication and fieldwork have not been performed. No drawing is actually AFC-approved.')
P('2. Planning basis','Heading1')
for t in ['Assumed NOA1November2026; actual award/start unconfirmed. Reviewed contract basis12months; conservative scenario last workday31October2027.','Two fabrication lines, one strengthening cell, three strengthening crews plus installation/support capacity, one heavy-lift and one cargo spread. Commitments, QA/coating/POB capacity and demonstrated productivity are not established.','Elapsed calendar-day onshore planning; no productive offshore work16May–15October. Each split activity must fit an open season. Safe stoppoints and renewed permits require engineering.','First-kit release is at20days inside a single35dayproduction stream. Duration allocations are illustrative, not vendor promises.','Forward/backwardpasses and named-exclusive-resource reservations are checked. This is not full labour/quantity-loaded P6 validation.']:P(t,'SmallP')
P('3. Quantity audit','Heading1')
T(['Item','Row sum','Printed total','Action'],[['Structural steel','1141MT','1142MT','Reconcile1MT'],['Handrail','1166m','1168m','Reconcile2m'],['Grating','718m2','719m2','Reconcile1m2'],['Demolition','357MT','357MT','Match arithmetic; actual scope still verify'],['GI sheeting','1274m2','1274m2','Match arithmetic; final detail verify']],[120,100,100,190])
P('These differences were verified against the uploaded MTO page1. Do not silently amend source quantities or place orders from unreconciled summary totals. Addendum wrap steel estimates are not automatically extra to these totals.')
P('4. Critical path and milestones','Heading1')
P('Selected driving route: kick-off → survey → early material release → procurement → BHS fabrication → BHS transport → BHS installation and service acceptance → oldbridge demolition → ICP–ICG installation → residual removal and offshore punch → dossiers → review/acceptance.53nonzeroactivities have zero float to internal earliest finish. Tied branches are retained in the detailed register.')
ms=['F020','I020','D020','I030','M100','O100','M200'];cb={x['id']:x for x in D['comparison']}
T(['Milestone boundary','3crew case','2crew case'],[[by[i]['name'],by[i]['finish'],cb[i]['finish']] for i in ms],[290,110,110])
P('Milestones are date boundaries: offshore physical works finish1May, milestone2May; demob work12–14May, milestone15May; handover milestone15June. The source methodology and operational approvals determine actual completion criteria.')
P('5. Calendar sensitivity and decision','Heading1')
T(['Extra critical material days','Offshore milestone','Demob boundary','Overall milestone'],[[x['extra_material_days'],x['offshore_milestone'],x['demob_phase_complete'],x['overall_milestone']] for x in D['sensitivity']],[120,130,130,130])
P('This is a conservative calendar stress test, not a probabilistic forecast:10reserved days are modeled as a seasonalcapacityhold, and marine tasks cannot use the closed season. The earlyJune target is not robustly validated. Increase evidenced earlyseason margin, verify workable stopping/revalidation and demob strategy, and confirm actual lead times/resources before approving a baseline.')
P('The2crew case includes an assumed7daysecond-seasonremobilisation and forecasts19November2027, beyond the unextended scenario deadline. Actual capacity, cost, safe suspension and contractual relief remain unresolved.')
P('6. Required release decisions','Heading1')
for t in ['Close geometry/model/load/reaction/camber/support queries and complete actual structural design checking.','Confirm quantities, supplier lead times, yard/inspection/coating capacity, thirdcrew and marine commitments.','Price cost, cash, standby, remobilisation, bank margins, billing/collection timing and funding.','Obtain applicable ONGC engineering/operational approvals and preserve AFC status honestly.']:P(t,'SmallP')
P('7. All zero-float nonzero tasks','Heading1')
T(['ID','Package / activity','Start','Lastworkday','Days'],[[x['id'],x['package']+' / '+x['name'],x['start'],x['finish'],x['duration']] for x in S if x['critical'] and x['duration']],[55,245,80,80,50])
def footer(canv,doc):canv.setFont('Helvetica',8);canv.drawString(36,20,'SHIVANSH / ONGC BRRP | P5 | FOR REVIEW');canv.drawRightString(A4[0]-36,20,str(doc.page))
SimpleDocTemplate(str(O/'BRRP_P5_Critical_Path_and_Project_Review.pdf'),pagesize=A4,leftMargin=36,rightMargin=36,topMargin=36,bottomMargin=36).build(story,onFirstPage=footer,onLaterPages=footer)
# AFC review form is expressly NOT an approval or a newly fabricated drawing.
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10);doc.add_heading('AFC Preparation and Review Control Forms',0)
doc.add_paragraph('P5 | FOR REVIEW ONLY | No drawing in this form is approved for construction. Construction release remains deferred.')
for title,items in [('Document identification',['Proposed drawing ID: __________________','Title and bridge: __________________','Revision / current issue purpose: __________________','Source survey and drawing revisions: __________________']),('Technical completeness',['Checked calculation references: __________________','Model/member/connection/bearing/temporary-work checks: __________________','Interface / service / marine reviews closed: __________________','Material specifications / MTO / inspection references: __________________']),('Review and release',['Prepared by / date: __________________','Checked by / date: __________________','Technical review comments and closure record: __________________','Applicable ONGC review / approval evidence: __________________','Authorised release/transmittal reference: __________________','STATUS: NOT AFC — HOLD until required evidence is recorded'])]:
 doc.add_heading(title,1)
 for x in items:doc.add_paragraph(x)
doc.add_page_break();doc.add_heading('Registered engineering deliverables',1)
t=doc.add_table(rows=1,cols=4);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Proposed ID','Bridge','Deliverable','Current status']):c.text=v
for x in G['deliverables']:
 for c,v in zip(t.add_row().cells,[x[0],x[1],x[2],'NOT AFC / ENGINEERING OPEN']):c.text=v
for row in t.rows:
 for cell in row.cells:
  for pp in cell.paragraphs:
   for run in pp.runs:run.font.size=Pt(8)
doc.save(O/'BRRP_P5_AFC_Preparation_Review_Forms.docx')
# Carry forward actual development evidence, not approval claims.
eng=O/'Engineering_Development_P4';eng.mkdir(exist_ok=True)
for n in ['BRRP_P4_Design_Development_Report.docx','BRRP_P4_Design_Development_Calculations.xlsx','BRRP_P4_Drawings_With_Tender_Landing_Appendix.pdf','BRRP_P4_Calculation_Snapshot.csv']:
 shutil.copyfile(R/'outputs/BRRP_Design_Development_P4'/n,eng/n)
# Clear manifest and deliverable status, with workbook as source for detailed data.
readme=['SHIVANSH / ONGC BRRP — INTEGRATED PROJECT PACKAGE P5','Prepared25October2026. FOR REVIEW. Construction release deferred.','', 'MAIN CONTENTS','1 Integrated Project Execution Plan (Word):scope,engineering,MTO,procurement,fabrication,marine/execution,commissioning,quality/HSE,cost,handover and detailed appendices.','2 Integrated Project Control (Excel):21sheets including205activityscenario,207activitytwo-crewcase,CPM/float,quantityreconciliation,100workpackages,101deliverables,25procurementpackages.','3 Overall Project Schedule (PDF).','4 Critical Path and Project Review (PDF).','5 AFC Preparation Review Forms (Word) — NOT approved drawings.','6 CSV schedule exports and transparent data JSON.','7 Engineering Development P4 folder:source-baseddrawings,calculationworkbook,reviewreport.','', 'STATUS','New global/member/connection/bearing/fatigue/liftingdesigns are not independently completed.','No AFC approval,purchase,charter,fabrication,installation,commissioning or clientacceptance is claimed.','All dates/resources/durations are planningassumptions pendingvalidation. ActualNOA is unconfirmed.','Source MTO row discrepancies:steel1141vs1142MT;handrail1166vs1168m;grating718vs719m2.','Primary milestone15June2027;two-crew comparator19November2027 withassumedsecond-seasonremobilisation.','The primarytarget is sensitive toseasonalclosure. See calendarstress tests beforebaselineapproval.','', 'FILES MAY BE DOWNLOADED LATER FROM THE WORKSPACE downloads FOLDER. Download support is outside these files.']
(O/'README_P5.txt').write_text('\n'.join(readme))
# Render previews for inspection.
for filename,n in [('BRRP_P5_Overall_Project_Schedule.pdf',0),('BRRP_P5_Critical_Path_and_Project_Review.pdf',0)]:
 p=fitz.open(O/filename);p[n].get_pixmap(matrix=fitz.Matrix(1.1,1.1),alpha=False).save(O/(Path(filename).stem+'_preview.png'));print(filename,'pages',len(p))
print('Prepared overall Gantt, criticalpathreviewPDF, AFCcontrolforms and carried-forwardP4evidence.')
