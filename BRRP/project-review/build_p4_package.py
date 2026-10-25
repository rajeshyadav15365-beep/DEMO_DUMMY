from pathlib import Path
import math,csv,json,shutil,zipfile
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
from docx import Document
from docx.shared import Pt,Inches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3,landscape
from reportlab.lib.colors import HexColor,black,white
import fitz
O=Path('/workspace/project-review/outputs/BRRP_Design_Development_P4');R=Path('/workspace/project-review');OLD=R/'outputs/bridge_engineering'
B=[dict(code='101',name='ICW-ICG LD',L=56,gap=54.5,b=1.2,left='ICW',right='ICG',report='BPR-0004 Rev0',pages='133-135;152-153',dv=105,dh=87,uc=.557,juc=.509,dl=85,LL=156.25,outLL=144.69,dead=418.03,ng=155.75,cable=.750,mon=0,pipe=0,source='ICW grid4 northside / ICG grid3 southside'),dict(code='102',name='ICP-ICG LD',L=84,gap=80.5,b=1.2,left='ICP',right='ICG',report='BPR-0003 Rev0',pages='196-199;207-209',dv=177,dh=123,uc=.610,juc=.625,dl=143,LL=288,outLL=288,dead=966.64,ng=156.73,cable=.882,mon=0,pipe=0,source='ICP grid3 northside / ICG scramble-net area'),dict(code='103',name='BHS-SLQ LD',L=42,gap=40,b=2.324,left='BHS',right='SLQ',report='BPR-0005 RevA',pages='97-100;122-123',dv=32,dh=76,uc=.916,juc=.859,dl=32,LL=268.42,outLL=268.42,dead=789.79,ng=266.47,cable=77.66,mon=20,pipe=1440,source='BHS nearA1 southside / SLQ nearB2 southside')]
# Fresh calculations only; no new global structural solve is claimed.
for b in B:
 b.update(q=250*9.81/1000,qlocal=500*9.81/1000,live_rect=b['L']*b['b']*250*9.81/1000,live_per_m=b['b']*250*9.81/1000,vlimit=b['L']*1000/400,hlimit=b['L']*1000/500,thermal_per_C=12e-6*b['L']*1000)
 b['vutil']=b['dv']/b['vlimit'];b['hutil']=b['dh']/b['hlimit']
sections=[]
for D,t,l in [(406,15.9,4),(406,12.7,4),(168,9.5,4.98),(114,7.9,4.71)]:
 d=D-2*t;A=math.pi*(D*D-d*d)/4;I=math.pi*(D**4-d**4)/64
 sections.append([D,t,d,A,I,2*I/D,(I/A)**.5,A*.00785,l,'Source ICP VIV table only; member mapping, resistance and selection NOT established'])
wraps=[]
for name,D,t0,L,t in [('ICG',660,12.7,15,10),('ICP',762,16,13.5,12),('BHS',457,12.7,3,10)]:
 vol=math.pi*((D+2*t)**2-D**2)/4*L/1e6;mass=vol*7850
 wraps.append([name,D,t0,L,t,math.pi*(D+t),vol,mass,mass/1000,'Ideal external full annulus only; four-quarter cuts, fit-up, corrosion, welds, end details, tolerances and wastage excluded. Not a purchase quantity.'])
criteria=[['Design life',25,'years','BRRP PartII p119','CONFIRMED'],['Elastic modulus E',210000,'MPa','PartI Annexure2 p66','CONFIRMED'],['Shear modulus G',80770,'MPa','PartI Annexure2 p66','CONFIRMED'],['Poisson ratio',.3,'ratio','PartI Annexure2 p66','CONFIRMED'],['Steel density',7850,'kg/m3','PartI Annexure2 p66','CONFIRMED'],['Gravity',9.81,'m/s2','PartI Annexure2 p66','CONFIRMED'],['Thermal expansion',.000012,'per C','PartI Annexure2 p66','CONFIRMED'],['Global walkway loading',250,'kg/m2','Bridge clause p42; Annexure4','CONFIRMED'],['Local grating loading',500,'kg/m2','Annexure4 p70','CONFIRMED'],['Minimum clear walkway',1.2,'m','Bridge clause p43','CONFIRMED'],['Bridge vertical limit','L/400','mm consistent with L','p43;Annexure12 pp85-86','CONFIRMED'],['Bridge lateral wind limit','L/500','mm consistent with L','Annexure12 p86','CONFIRMED'],['Differential translation multiplier',1.5,'ratio','Bridge clause p43','CONFIRMED; input movements pending'],['Hinged axial thrust multiplier',1.5,'ratio','Bridge clause p43','CONFIRMED; thrust envelope pending'],['Future piping operating-load allowance',20,'percent','Bridge clause p42','CONFIRMED; included/excluded basis to reconcile'],['Camber','Compensate dead-load deflection','profile','Bridge clause p43','RULE CONFIRMED; magnitude/profile pending'],['Wind gust for56m/84m',15,'seconds','Annexure5 p73','CONFIRMED'],['Wind gust for42m',3,'seconds','Annexure5 p73','CONFIRMED'],['Environment','MH North and South','field','PartII p119;Annexure13A','CONFIRMED; height adjustment/calculation pending'],['Material procurement','Spec6001F Rev14','requirements','Specification Combined pp169-243','Grade/toughness selection pending'],['Welding/NDT','FS2009F Rev8','requirements','Specification Combined pp3-38','Apply by joint and inspection plan'],['Coating','FS2005 Rev9','requirements','Specification Combined pp91-168','Exposure-specific system selection pending'],['Roof','Process-platform connecting bridges; slope minimum1:50','applicability','Bridge clause p43','Confirm bridge-specific implementation'],['Location tolerance','+1.0m as written in source','requirement','Bridge clause p43','Clarify interpretation for existing supports; NOT bearing travel']]
loads=[['LC01','In-place extreme','Bridge dead + services dead + piping operating contents + extreme wind + friction','Include applicable allowances; verify source load map'],['LC02','In-place operating','Bridge dead + services dead + piping contents + walkway LL + monorail LL + operating wind + friction','Assess both wind signs/directions and service configurations'],['LC03','Dead / dry reference','Bridge dead + piping/cable/tray dead','Camber and construction-setting reference must be defined'],['LC04','Hydrotest','Bridge dead + services dead + one large pipe hydrotest at a time','Check actual line/test sequence; not all lines assumed simultaneous'],['LC05','Movement / supports','Relative platform translations and compatible rotations','150% specified translation;150% expected hinged axial thrust; no double counting'],['LC06','VIV / fatigue','Wind-induced vibration of truss members','Displacement and fatigue damage; in-service and applicable transport states'],['LC07','Fabrication / weighing','Actual support points and construction loads','Check changed support configurations and local effects'],['LC08','Loadout','Actual loadout support positions / loss of support / yard bearing','Procedure-specific; applicable three-point cases'],['LC09','Transport','Actual barge/vessel motions, cargo mass/COG and restraints','Criteria preliminary motions are screening basis; final route/spread analysis required'],['LC10','Lifting / landing','Rigging geometry, skew distribution, weight/COG and dynamics','Operation-specific 3D model; no approved lift points defined'],['LC11','Local grating / fittings','Local element loading and appropriate concentrated loads','500kg/m2 grating local check distinct from250kg/m2 global walkway'],['LC12','Landings / jackets','Concurrent signed bridge reactions plus existing platform loads','Incorporate addendum strengthening and actual condition']]
queries=[['Q01','Final bearing centres, levels, skew and condition','OPEN','Survey-dependent geometry; model length is not shop length'],['Q02','Native analysis connectivity/member groups/load map','OPEN','No independently reproduced FE model in this package'],['Q03','Actual live-load areas and report/model differences','OPEN','Rectangle screens do not replace documented load generation'],['Q04','Camber instruction in reports versus criteria rule','PARTLY RESOLVED','Dead-load compensation governs; profile and construction datum still required'],['Q05','BHS SLQ support wording','OPEN','Report and methodology support fixed intent; SOW conflicting wording needs formal resolution'],['Q06','Wind duration difference','EXPLAINED','15sec for>50m;3sec for<50m; height/speed derivation pending'],['Q07','Source member schedule / code editions / effective lengths','OPEN','VIV list is not a full member schedule'],['Q08','Signed concurrent support reactions / existing capacity','OPEN','Summary values cannot be assembled into a fictitious concurrent case'],['Q09','Missing101/102 main GA and0005;103series drawings','OPEN','Existing drawings and source studies used; gaps remain recorded'],['Q10','Material grade, toughness and corrosion application','OPEN','50ksi yield in study does not establish full material procurement specification'],['Q11','BHS piping20%future allowance','OPEN','Confirm whether source1440kN already contains it; avoid duplication'],['Q12','Weights and lift modules','OPEN','Tender MT values are not verified lifted masses'],['Q13','MTO1MT difference','OPEN','Existing landing strengthening rows216vs217MT summary'],['Q14','Other-project documents','CONTROLLED','BRRP PartII adopted; S&SRP/WSSRP-2 project-specific scope excluded unless directed']]
# Workbook: static audit is available even without Excel recalculation.
wb=openpyxl.Workbook();wb.remove(wb.active)
def tab(name,head,rows,widths=None):
 s=wb.create_sheet(name);s.append(head)
 for r in rows:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='17365D')
 for i in range(1,s.max_column+1):s.column_dimensions[openpyxl.utils.get_column_letter(i)].width=widths[i-1] if widths and len(widths)>=i else 24
 for row in s:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
 return s
read=[['Status','P4 PRELIMINARY DESIGN DEVELOPMENT. Prepared for review; construction release deferred by user.'],['Executed work','Source criteria consolidation; independent arithmetic and section geometry; thermal sensitivity; annular wrap steel estimate; source-based drawing and interface detailing.'],['Not executed','No independent 3D FE solve, final member capacity design, fatigue model, connection resistance, bearing selection or engineered lifting release.'],['Input hierarchy','BRRP PartII + general PartI Rev14 + structural scope/addendum + applicable specifications. Resolve conflicts through project technical queries.'],['Drawing dimensions','56/84/42m and walkway loading widths are study references; source drawings retain tender status.'],['Static vs formulas','Audit Snapshot and report contain evaluated numeric results. Formula sheets recalculate in Excel; blank project inputs produce INPUT REQUIRED.'],['Source results','All UC/deflection outputs are attributed to earlier reports, not recalculated structural analysis.'],['Thermal examples','10C and20C are sensitivity increments, not adopted site design temperatures.'],['Wrap estimate','Geometric shell steel only; no construction or procurement release.'],['Acceptance','No claim of structural adequacy, approval or certification is made by this preliminary package.']]
tab('READ ME',['Topic','Basis'],read,[28,135])
tab('Design Criteria',['Parameter','Value','Unit / interpretation','Source','Status'],criteria,[40,65,34,48,64])
tab('Bridge Inputs',['Code','Bridge','Model L m','Approx gap m','Loading width m','Sliding end','Fixed intent','Source report','PDF pages'],[[b['code'],b['name'],b['L'],b['gap'],b['b'],b['left'],b['right'],b['report'],b['pages']] for b in B])
tab('Source Load Register',['Bridge','Report DL kN','Report NGDL kN','Cable load kN','Monorail LL kN','Piping operating kN','Report walkway LL kN','Output walkway LL kN','Status'],[[b['name'],b['dead'],b['ng'],b['cable'],b['mon'],b['pipe'],b['LL'],b['outLL'],'Source table only; no load combination/allowance adoption'] for b in B])
tab('Load Cases',['ID','Condition','Required assembly','Review action'],loads,[14,30,105,95])
calc=[]
for i,b in enumerate(B,2):calc.append([b['name'],b['L'],b['b'],250,9.81,f'=D{i}*E{i}/1000',f'=C{i}*F{i}',f'=B{i}*G{i}',b['LL'],f'=I{i}-H{i}',f'=B{i}*1000/400',b['dv'],f'=L{i}/K{i}',f'=B{i}*1000/500',b['dh'],f'=O{i}/N{i}'])
tab('Calculation Formulas',['Bridge','Study L m','Loading width m','Mass LL kg/m2','g m/s2','q kN/m2','Line LL kN/m','Rectangle LL kN','Source LL kN','Source minus rectangle kN','Vertical limit mm','Source vertical mm','Vertical ratio','Lateral limit mm','Source lateral mm','Lateral ratio'],calc)
tab('Audit Snapshot',['Bridge','q global kN/m2','q grating local kN/m2','Line LL kN/m','Rectangle LL kN','Source LL kN','Vertical limit mm','Source vertical mm','Ratio vertical','Lateral limit mm','Source lateral mm','Ratio lateral','Listed member UC','Listed joint UC'],[[b['name'],b['q'],b['qlocal'],b['live_per_m'],b['live_rect'],b['LL'],b['vlimit'],b['dv'],b['vutil'],b['hlimit'],b['dh'],b['hutil'],b['uc'],b['juc']] for b in B])
tab('Section Geometry',['Source OD mm','Source wall mm','ID mm','Area mm2','I mm4','Elastic Z mm3','Radius gyration mm','Mass kg/m','Source VIV length m','Adoption status'],sections)
tab('Thermal Sensitivity',['Bridge','L study m','Alpha per C','Expansion mm per C','At10C mm','At20C mm','Status'],[[b['name'],b['L'],.000012,b['thermal_per_C'],b['thermal_per_C']*10,b['thermal_per_C']*20,'Sensitivity only; design thermal range and movement combination pending'] for b in B])
mov=[]
for i,b in enumerate(B,2):mov.append([b['name'],None,None,f'=IF(COUNT(B{i}:C{i})<2,"INPUT REQUIRED",1.5*(ABS(B{i})+ABS(C{i})))',None,f'=IF(COUNT(E{i})<1,"INPUT REQUIRED",1.5*ABS(E{i}))','Enter directional absolute platform movements and compatible thrust; rotation, thermal and guides require separate compatible checks'])
tab('Movement Inputs',['Bridge','Platform A translation mm','Platform B translation mm','150% worst-sum mm','Expected axial thrust kN','150% thrust kN','Application'],mov)
tab('Wrap Steel Estimate',['Platform','Existing OD mm','Existing nominal wall mm','Length m','Wrap wall mm','Neutral circumference mm','Ideal shell volume m3','Ideal shell mass kg','Ideal shell mass MT','Limitations'],wraps)
tab('Open Items',['ID','Subject','Status','Required action'],queries,[14,65,28,110])
wb.calculation=openpyxl.workbook.properties.CalcProperties(fullCalcOnLoad=True)
wb.save(O/'BRRP_P4_Design_Development_Calculations.xlsx')
# static tables for portable review
with open(O/'BRRP_P4_Calculation_Snapshot.csv','w',newline='') as f:
 w=csv.writer(f);w.writerow(['Bridge','ModelL_m','Walkway_q_kPa','LineLL_kNperm','RectangleLL_kN','VerticalLimit_mm','SourceVertical_mm','LateralLimit_mm','SourceLateral_mm','Thermal_mm_per_C']);w.writerows([[b['name'],b['L'],b['q'],b['live_per_m'],b['live_rect'],b['vlimit'],b['dv'],b['hlimit'],b['dh'],b['thermal_per_C']] for b in B])
# Word calculation report.
d=Document();sec=d.sections[0];sec.header.paragraphs[0].text='SHIVANSH / ONGC BRRP | P4 DESIGN DEVELOPMENT | FOR REVIEW'
d.styles['Normal'].font.name='Calibri';d.styles['Normal'].font.size=Pt(10)
d.add_heading('Three Replacement Bridges — Design Development',0);d.add_paragraph('P4 | 25 October 2026 | Construction release deferred')
d.add_paragraph('This package develops the supplied structural criteria and replacement studies into a traceable preliminary engineering basis, calculation sheets and source-based detailing. It does not constitute completed member, connection, bearing, fatigue or marine-operation design. Reported structural analysis results remain attributed to the source studies.')
def table(head,rows):
 t=d.add_table(rows=1,cols=len(head));t.style='Light Shading Accent 1'
 for c,v in zip(t.rows[0].cells,head):c.text=str(v)
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=str(v)
 return t

d.add_heading('1. Scope and source hierarchy',1)
d.add_paragraph('Scope: ICW–ICG LD, ICP–ICG LD and BHS–SLQ replacement/new bridges, with landings, platform/jacket interfaces, access fittings and service supports. BHS includes service-transfer and demolition interfaces. BRRP-specific PartII is the project basis; other-project PartII content in Section3 is not imported.')
d.add_paragraph('Principal sources: Design Criteria & Execution Methodology.pdf (PartI Rev14 pp1–116; BRRP PartII pp117–120; PartIII pp121–132; methodology RevC pp133–154); SOW & Bridge Report Combined.pdf (structural scope pp15–53 and reports as below); MTO Drawings IRT.pdf; structural addendum; Specification Combined.pdf. Source IDs, revisions and applicable clauses govern; conflicts remain queries.')
table(['Bridge','Study L / approximate gap','Sliding / fixed intent','Study report / combined PDF pages'],[[b['name'],f'{b["L"]} / {b["gap"]} m',b['left']+' / '+b['right'],b['report']+' / '+b['pages']] for b in B])
d.add_heading('2. Consolidated design basis',1)
for row in criteria:d.add_paragraph(f'{row[0]}: {row[1]} {row[2]}. Source: {row[3]}. {row[4]}.')
d.add_heading('3. Load model preparation',1)
d.add_paragraph('Separate structural self-weight, non-generated dead loads, services, walkway loading, local grating loading, operating contents and future allowances. Record source revision, included contingencies and area/length distribution. Source report totals are not directly transferred to a revised geometry without reconciliation.')
table(['Bridge','Report DL','NGDL','Cable','Walkway LL','Piping operating'],[[b['name'],b['dead'],b['ng'],b['cable'],b['LL'],b['pipe']] for b in B]);d.add_paragraph('Table units kN; transcribed from source load tables. Source tables and generated output differ in some entries; these are not final design totals.')
for x in loads:d.add_paragraph(f'{x[0]} {x[1]}: {x[2]}. {x[3]}.')
d.add_paragraph('BHS future-piping sensitivity: if the reported1440kN represents the same operating-load basis before the20% future allowance, the corresponding value is1728kN. Do not apply this increment until existing allowance treatment is confirmed. Operating contents, empty pipe weight and local support actions must remain consistent.')
d.add_heading('4. New arithmetic: walkway and local grating loads',1)
d.add_paragraph('Using criteria g=9.81m/s2: global walkway q=250x9.81/1000=2.4525kN/m2; local grating q=500x9.81/1000=4.905kN/m2. Line load is q times actual loaded width. Rectangle load is q x studyL x report loading width; this is a reconciliation screen, not a substitute for the real load map.')
table(['Bridge','Line load kN/m','Rectangle load kN','Source walkway LL kN'],[[b['name'],f'{b["live_per_m"]:.4f}',f'{b["live_rect"]:.3f}',b['LL']] for b in B])
d.add_paragraph('ICW source walkway table156.25kN differs from output144.69kN. ICP source288kN and BHS268.42kN do not equal the simple rectangles shown. Confirm actual areas, extensions, revisions and generation rules before final design.')
d.add_heading('5. Source serviceability and utilisation audit',1)
table(['Bridge','Vertical / limit mm','Ratio','Lateral / limit mm','Ratio','Listed member / joint UC'],[[b['name'],f'{b["dv"]}/{b["vlimit"]:.0f}',f'{b["vutil"]:.1%}',f'{b["dh"]}/{b["hlimit"]:.0f}',f'{b["hutil"]:.1%}',f'{b["uc"]}/{b["juc"]}'] for b in B])
d.add_paragraph('The limits are independently recalculated from the criteria; deflections and UC values are source results, not results of a new model. Highest listed member/joint values are not independently verified full envelopes. No adequacy statement is made for changed loads or geometry. Do not infer allowable load increases from these ratios.')
d.add_heading('6. Bearing movements, thermal sensitivity and camber',1)
d.add_paragraph('Use directional differential platform movements and compatible rotations. Criteria translation demand is150% of the worst-case sum of absolute adjacent-platform translations. Hinged axial connection demand is150% of expected axial thrust. Bearing geometry, lateral/vertical guides, friction, uplift, travel, maintenance and tolerances require a checked support model. The location-tolerance clause is not an instruction to specify1m bearing travel.')
d.add_paragraph('Thermal expansion sensitivity: deltaL=alpha x L x deltaT; alpha=12x10^-6/C.10C and20C below are illustrative increments, not adopted site temperature ranges. Combine thermal and imposed movement compatibly; avoid counting the same movement twice.')
table(['Bridge','mm per C','At10C mm','At20C mm'],[[b['name'],f'{b["thermal_per_C"]:.3f}',f'{b["thermal_per_C"]*10:.2f}',f'{b["thermal_per_C"]*20:.2f}'] for b in B])
d.add_paragraph('Camber must compensate dead-load deflection under the approved geometry and construction-setting assumptions. Source DL deflections are85,143 and32mm respectively. These values are references, not adopted fabrication camber. The source-study1:200 statement requires reconciliation; no L/200 camber is issued.')
d.add_heading('7. Source section geometry',1)
d.add_paragraph('The ICP VIV table lists CHS406x15.9,406x12.7,168x9.5 and114x7.9mm. Workbook calculates ideal uncorroded A,I,Z,r and mass. It does not assign these to chords/braces, select sections, establish effective lengths or calculate code resistance.')
table(['OD x t mm','Area mm2','I mm4','Mass kg/m'],[[f'{r[0]}x{r[1]}',f'{r[3]:.1f}',f'{r[4]:.3e}',f'{r[7]:.2f}'] for r in sections])
d.add_heading('8. Addendum jacket wrap steel: preliminary quantity calculation',1)
d.add_paragraph('For a close-fitting external wrap idealised as a full annulus: V=pi/4[(D+2t)^2-D^2]L, with consistent units; mass=Vx7850. Four-quarter fabrication, fit-up gaps, existing corrosion, welds, ends, allowances and local attachments require detailing. This is only geometric steel accounting, not strengthening-capacity verification or a purchase quantity.')
table(['Platform','Existing OD x wall mm','Wrap length m','Wrap t mm','Ideal shell mass MT'],[[r[0],f'{r[1]}x{r[2]}',r[3],r[4],f'{r[8]:.3f}'] for r in wraps])
d.add_heading('9. Source-based detailing approach',1)
d.add_paragraph('Nine P4 development sheets cover three source-informed arrangements, three support/interface schematics and three bridge-specific fabrication/detailing schedules. Six original landing sheets are appended unchanged as tender-source references. New IDs are Shivansh development IDs and are not substitutions for EIL drawing numbers.')
d.add_paragraph('Dimensions and source sections are not inferred from screenshot scale. Connections, welds, bolt sizes, bearing sizes, guide gaps, lift points, member thicknesses and final camber are left explicitly unselected pending corresponding calculations. This preserves a usable preliminary package without fabricating detailed-design evidence.')
d.add_heading('10. Model reconstruction and engineering continuation',1)
for t in ['Verify actual nodes, axes, member connectivity, offsets, end releases, section groups and support freedoms.','Reproduce source gravity loadcases first and reconcile applied load totals, reactions and deflections.','Apply confirmed bridge-specific wind duration, environmental adjustment, service loads and allowances.','Check members for applicable strength, buckling, slenderness, fatigue and serviceability; check joints locally.','Transfer concurrent signed reaction envelopes to landing/deck/jacket models and incorporate addendum strengthening.','Develop operation-specific loadout, transport, lifting and temporary-condition analyses.','Issue preliminary framing/connection/bearing schedules with calculation references; resolve review comments before any later construction release.']:
 d.add_paragraph(t,style='List Bullet')
d.add_heading('11. Open-item register',1)
for q in queries:d.add_paragraph(' | '.join(q))
d.add_heading('12. Review status',1)
d.add_paragraph('Completed: criteria consolidation, source/assumption traceability, load reconciliation arithmetic, source serviceability arithmetic, thermal sensitivities, source section geometry, ideal addendum wrap quantities, source-based drawing sheets and input/review registers. Outstanding: independent global model and resistance design, survey verification, full member schedule, reaction envelopes, permanent and temporary connection/bearing/lifting calculations. Construction release remains deferred.')
d.save(O/'BRRP_P4_Design_Development_Report.docx')
# Portable JSON for exact audit and drawing generation.
(O/'P4_calculation_data.json').write_text(json.dumps({'bridges':B,'sections':sections,'wraps':wraps,'criteria':criteria,'loadcases':loads,'queries':queries},indent=2))
# Meaningful arithmetic checks, dimensional closure and file integrity.
assert abs(B[0]['live_rect']-164.808)<1e-9
assert abs(B[1]['live_rect']-247.212)<1e-9
assert abs(B[2]['hutil']-76/84)<1e-12
assert all(r[2]>0 and r[3]>0 and r[4]>0 for r in sections)
for r in wraps:
 mean_area=math.pi*(r[1]+r[4])*r[4]
 assert abs(r[6]-mean_area*r[3]/1e6)<1e-10
for p in O.iterdir():
 if p.suffix in ['.docx','.xlsx']:
  with zipfile.ZipFile(p) as z:assert z.testzip() is None
print('Prepared P4 report,12-sheet calculation workbook, static calculation CSV and audit data. Independent arithmetic checks passed.')
for r in wraps:print('Ideal wrap',r[0],round(r[8],4),'MT')
