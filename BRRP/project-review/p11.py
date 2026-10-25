import json,pathlib,hashlib,fitz,openpyxl
from docx import Document
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P11_Tender_Review');p.mkdir(exist_ok=True);rs=json.loads(pathlib.Path('/workspace/project-review/p11_manifest.json').read_text());current={r['path']for r in rs};prior={}
for f in pathlib.Path('/workspace/attachments').rglob('*'):
 if f.is_file() and str(f)not in current and f.suffix.lower() in ['.pdf','.docx','.xlsx']:prior.setdefault(hashlib.sha256(f.read_bytes()).hexdigest(),[]).append(str(f))
seen={};es=[]
for r in rs:
 f=pathlib.Path(r['path']);h=hashlib.sha256(f.read_bytes()).hexdigest();r.update(name=f.name,sha256=h,prior=prior.get(h,[]),duplicate=seen.get(h));seen.setdefault(h,f.name)
 if r['duplicate']:continue
 e={'name':f.name}
 if f.suffix=='.pdf':
  d=fitz.open(f);e['pages']=[{'page':i+1,'text':' '.join(pg.get_text().split())}for i,pg in enumerate(d)];r['pages']=len(d)
 elif f.suffix=='.docx':
  d=Document(f);e['paragraphs']=[{'paragraph':i+1,'text':x.text}for i,x in enumerate(d.paragraphs)if x.text.strip()];e['tables']=[[[c.text for c in row.cells]for row in t.rows]for t in d.tables]
 else:
  w=openpyxl.load_workbook(f,data_only=False);e['sheets']=[{'name':s.title,'rows':s.max_row,'cols':s.max_column,'cells':{c.coordinate:str(c.value)for row in s for c in row if c.value is not None}}for s in w]
 es.append(e)
(p/'inventory.json').write_text(json.dumps(rs,indent=2));(p/'extracts.json').write_text(json.dumps(es,indent=2))
print(json.dumps([{'name':r['name'],'prior':len(r['prior']),'duplicate':r['duplicate'],'pages':r.get('pages')}for r in rs],indent=2))
for e in es:
 if e['name'].startswith(('TBO','Technical_SCC')):
  print('\nFILE',e['name'])
  for pg in e['pages']:print('PAGE',pg['page'],pg['text'])
