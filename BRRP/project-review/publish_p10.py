import json,pathlib,shutil
from docx import Document
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Alignment,Font
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P10_Input_Review');inv=json.loads((p/'inventory.json').read_text());es=json.loads((p/'extracts.json').read_text());issues=json.loads((p/'csv_logic_exceptions.json').read_text())
findings=[
('Scope applicability','Section1.pdf pp1,6-8; Section2.pdf p1; Section3.pdf p1','All three Sections identify S&SRP (Sump Caisson and SWLP Casing Replacement Project). Section1 describes SWLP casings, vent booms, heaters, PBIV, caissons and RS-21 repairs. Shared platform names do not prove BRRP scope.','Prepare an applicability matrix linking any adopted requirement to BRRP incorporation/contract reference. Retain these sources separately until applicability is established.'),
('New source contents','inventory.json','Section1, the CSV, Stage4 and Stage5 are four distinct contents without exact prior local matches. Ten successfully transferred uploads match earlier local files; Stage10 variants are also identical to one another.','Do not assume no prior hash match proves a later revision; reconcile cell-level changes and issuance before superseding earlier plans.'),
('Network logic','CSV BR0002, BR0003, BR0005, BR0007, BR0008; complete exceptions sheet','87 activity records. Screening finds54 FS-linked successors starting on or before their predecessor finish day, using the supplied day numbers and assuming no lag because no lag field is present. BR0002 starts day5 after BR0001 finishing day10; BR0007 starts76 after BR0006 finishing99. Same-day cases depend on timing convention.','Resolve intended FS/SS relations, lags and inclusive-day conventions; validate calendars and import mapping. Do not call these dates a calculated CPM baseline or change relations merely to preserve target dates.'),
('Resource feasibility','Stage4 Resource Estimates C2:I11; Stage5 Yard Production C2:J11','Productivity and crew/vessel inputs are unfilled in the sampled resource tables. Stage5 yard allocations are TBD and release gates pending; Ready for loadout is NO.','Obtain yard outputs, working shifts, named crews, crane charts, vessel availability and shutdown windows; then calculate durations and resource-level the network.'),
('Spreadsheet calculation state','Workbook sheet inventory','Some Stage4/5 formulas have no stored calculated value; many deliberately return blank for missing inputs. Stage25 also has uncached formula outputs. This does not establish an Excel error; Python inspection does not recalculate formulas.','Distinguish missing input, blank formula result and stale/missing cache. Recalculate an editable copy in Excel after verified inputs and perform checks; no recalculation claim is made here.'),
('Cost estimate completeness','Stage25 READ ME B2:B9; Stage26 Commercial Inputs B3:B13 and Controls and Caveats','Financial models explicitly remain illustrative/unpriced. Zero contract value and zero cost inputs mean unknown, not zero exposure. Tonnage allocation is provisional.','Build priced WBS estimates from quotations, quantities, marine spread days, overheads, risk and taxes; reconcile to signed terms and actual commercial status.'),
('Funding horizon','Stage26 Controls and Caveats rows5-9; CEO Decision Dashboard C5','29-week scenario excludes part of contract life. The workbook states retention is withheld but eventual release is not added; advance recovery needs verification; proposed certification/payment lags require confirmation.','Extend cash flow through completion, payment delays, retention/warranty/security release and final account. Add explicit advance recovery, BG margin release, tax and financing events and reconcile against milestones; then recalculate peak funding.'),
('Contract duration','Stage25 READ ME B3:B4; earlier Technical SCC completion provision','1 November2026 and360 days are planning assumptions. Earlier reviewed SCC states12 months from NOA;360 days is not automatically equivalent to12 months.','Use controlled governing SCC/award, actual commencement evidence, monsoon rules and approved extensions; align all workbooks to one date/calendar basis.'),
('MTO reconciliation','Stage25 READ ME B5:B7; Stage10 CEO Decisions row3','Printed structural total1142MT versus detailed rows1141MT remains acknowledged and unresolved. Labels such as Verified MTO are not approval evidence.','Maintain discrepancy register and obtain detailed engineering quantity reconciliation before procurement commitments; retain other prior handrail/grating discrepancies.'),
('Baseline gates','Level3 P6 Preparation, Baseline Readiness rows2-13; Stage10 360D Logic','Workbooks identify award, survey, lift/COG, vessels, shutdowns, productivity, DCI/MCI and CPM as open. Target-day gates are proposals.','Connect each gate to an evidence owner, due date, document and acceptance; do not promote planned day360 completion or dashboard labels to proof of readiness.'),
('Newer tender amendments','P9 impact report; C8 DCI/MCI previously reviewed','This upload does not itself demonstrate integration of revised security, steel-origin, vendor12.3, review-working-day, welding/coating and PPE requirements.','Cross-reference P9 findings into procurement, engineering review, QA, financing and schedule predecessors before the next consolidated baseline issue.'),
('Unreviewed large PDF','SHIVANSH ONGC Offshore Bridges - PRELIMINARY.pdf','Transfer failed: file exceeds the executor32MiB limit. Its content was not opened or reviewed. The similarly named Excel was opened but equivalence cannot be assumed.','Keep PDF explicitly pending. The other14 successfully transferred files have been inventoried without requiring this PDF.')]
w=Workbook();s=w.active;s.title='Findings';s.append(['Topic','Source','Finding','Required action'])
for r in findings:s.append(r)
s=w.create_sheet('Inventory');s.append(['File','Bytes','SHA256','Prior exact matches','Batch duplicate','PDF pages'])
for r in inv:s.append([r['name'],r['result']['size_bytes'],r['sha256'],len(r['prior']),r['duplicate'],r.get('pages')])
s.append(['SHIVANSH ONGC Offshore Bridges - PRELIMINARY.pdf','TRANSFER FAILED: over32MiB; not reviewed'])
s=w.create_sheet('FS Date Exceptions');s.append(['Activity','Start day','Predecessor','Predecessor finish day'])
for r in issues:s.append(list(r.values()))
s=w.create_sheet('Workbook Sheet Audit');s.append(['File','Sheet','Rows','Columns','Formula count','No cached value'])
for e in es:
 for x in e.get('sheets',[]):s.append([e['name'],x['name'],x['rows'],x['cols'],x['formulas'],x['uncached']])
for s in w:
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.font=Font(bold=True)
 for col in s.columns:s.column_dimensions[col[0].column_letter].width=48
 for row in s:
  for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
w.save(p/'BRRP_P10_Input_and_Planning_Audit.xlsx')
d=Document();d.add_heading('BRRP — Additional input and planning-model review',0);d.add_paragraph('P10 |25 October2026 | Source screening and internal consistency audit')
d.add_paragraph('15 files supplied;14 transferred and opened;13 distinct accessible contents; four contents have no exact match among earlier local attachments. One PDF could not transfer because it exceeds32MiB. Section1/2/3 contain12/43/323 pages respectively. PDF title/scope screening, workbook data/formula inventory and CSV day-logic screening were performed. This is not an exhaustive numerical review of all378 PDF pages, formula recalculation, P6 scheduling run, structural analysis or AFC approval.')
for title,source,finding,action in findings:
 d.add_heading(title,1);d.add_paragraph('Source: '+source);d.add_paragraph(finding);d.add_paragraph('Required update: '+action)
d.add_heading('Current disposition',1);d.add_paragraph('Keep S&SRP applicability unresolved, repair and validate schedule logic, secure resource and commercial inputs, reconcile MTO and current amendments, extend financial horizons, then recalculate and obtain baseline approval. Preserve original uploaded files. This review has not changed source workbook formulas or overwritten the earlier planning scenario. Actual contract status remains unverified; text in the uploaded dashboards is source content, not independent confirmation.')
d.save(p/'BRRP_P10_Additional_Input_Review.docx')
for name in ['BRRP_P10_Additional_Input_Review.docx','BRRP_P10_Input_and_Planning_Audit.xlsx']:
 dest=pathlib.Path('/workspace/downloads')/name;shutil.copy2(p/name,dest);dest.chmod(0o644)
assert len(load_workbook(p/'BRRP_P10_Input_and_Planning_Audit.xlsx').worksheets)==4
print('Saved and reopened P10 report and four-sheet audit;54 FS date exceptions recorded.')
