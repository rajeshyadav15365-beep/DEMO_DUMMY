import pathlib,json,datetime as dt,collections,csv,hashlib
from openpyxl import load_workbook
from docx import Document
import fitz
p=pathlib.Path('/workspace/project-review/outputs/BRRP_P14_Consolidated_Management');checks=[]
for q in [p,p/'MORE_YARD',p/'FAST_SUPPLY_MORE_YARD']:
 d=json.loads((q/'P14_schedule_data.json').read_text());ts=d['tasks'];by={t['id']:t for t in ts};assert len(by)==285
 for t in ts:
  s=dt.date.fromisoformat(t['ES']);e=dt.date.fromisoformat(t['EF']);assert e>=s
  assert all(t['ES']>=by[x]['EF']for x in t['preds'])
  days=[s+dt.timedelta(days=i)for i in range((e-s).days)]
  if t['calendar']=='OFF':assert all(not dt.date(z.year,5,16)<=z<=dt.date(z.year,10,15)for z in days);count=len(days)
  elif t['calendar']=='OFFICE':count=sum(z.weekday()<5for z in days)
  elif t['calendar']=='YARD':count=sum(z.weekday()<6for z in days)
  else:count=len(days)
  assert count==t['duration'],(t['id'],count,t['duration'])
  assert t['total_float_calendar_days']>=0
 resources=collections.defaultdict(list)
 for t in ts:
  if t['resource']:resources[t['resource']].append(t)
 for r,ls in resources.items():
  ls.sort(key=lambda x:x['ES'])
  assert all(ls[i]['EF']<=ls[i+1]['ES']for i in range(len(ls)-1)),r
 used={x for t in ts for x in t['preds']};assert [t['id']for t in ts if not t['preds']]==['START'];assert [t['id']for t in ts if t['id']not in used]==['FINISH']
 checks.append({'scenario':d['scenario'],'activities':len(ts),'finish':d['finish'],'checks':'FS links, durations/calendar, offshore exclusion, named-resource nonoverlap, unique start/finish, nonnegative scenario float passed'})
b=load_workbook(p/'BRRP_P14_Integrated_Management_Control.xlsx');assert len(b.worksheets)==34
assert b['Procurement 25'].max_row==26;assert b['Engineering DCI'].max_row==184;assert b['Technical Queries'].max_row==16
assert b['156 Week Cash']['K2'].value=='=IF(OR(N2<>"VERIFIED",COUNT(C2:J2)<8),"",SUM(C2:E2)-SUM(F2:J2))'
# Structural formula check only: no spreadsheet calculation engine was used.
for s in b:
 for row in s:
  for c in row:
   if c.data_type=='f':assert '#REF!'not in c.value
Document(p/'BRRP_P14_Project_Execution_Management_Manual.docx')
for f in p.glob('*.pdf'):
 d=fitz.open(f);assert len(d)>0;assert all(len(pg.get_text())>50for pg in d)
r={'scenario_checks':checks,'workbook_sheets':len(b.worksheets),'spreadsheet_validation':'Files reopened; counts and references checked; Excel formulas not recalculated here','structural_design_validation':'Not performed; actual independently checked design and approval gates remain open','files':[{ 'file':f.name,'bytes':f.stat().st_size,'sha256':hashlib.sha256(f.read_bytes()).hexdigest()}for f in p.iterdir()if f.suffix in ['.docx','.pdf','.xlsx','.csv']]};(p/'P14_quality_checks.json').write_text(json.dumps(r,indent=2));print(json.dumps({'scenarios_verified':len(checks),'workbook_sheets':len(b.worksheets),'files_reopened':len(r['files']),'status':'passed within stated scope'}))
