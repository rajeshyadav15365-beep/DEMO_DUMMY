exec(open('/workspace/project-review/build_schedule.py').read().split('# Save data')[0])
import copy,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table,TableStyleInfo
from openpyxl.chart import BarChart,Reference
BASE=copy.deepcopy(A)
# Recovery: early phased strengthening kit release plus third offshore crew.
REC=copy.deepcopy(BASE)
ix=next(i for i,a in enumerate(REC) if a['id']=='F040')
first=copy.deepcopy(REC[ix]);first.update(id='F041',name='First approved strengthening kits: phased fabrication / QA release',dur=20,gate='Bridge-specific release required; remaining kits continue in same cell')
REC.insert(ix,first)
# F041 is a progress marker within F040, not an additional fabrication-cell demand.
first['resource']='Milestone within strengthening cell production'
for a in REC:
 if a['id']=='O010':a['preds']=['F041','MR10']
 if a['id']=='D030':a['preds']=list(dict.fromkeys(a['preds']+['RA03']))
 if a['id']=='RA04':a['preds']=['O010','F040'];a['resource']='Strengthening crew C (additional)'
 if a['id'] in ['RA01','RA02','RA03','RB01','RB02','RB03']:a['preds']=list(dict.fromkeys(a['preds']+['F040']))

def seasonal_start(es,dur):
 while dur:
  sd=START+timedelta(days=es);ed=sd+timedelta(days=dur-1)
  conflict=None
  for y in range(sd.year,ed.year+1):
   m0=date(y,5,16);m1=date(y,10,15)
   if sd<=m1 and ed>=m0:conflict=date(y,10,16);break
  if conflict is None:return es
  es=(conflict-START).days
 return es

def calculate(tasks):
 by={};
 for a in tasks:
  raw=max([0]+[by[p]['ef'] for p in a['preds']]);es=seasonal_start(raw,a['dur']) if a['off'] else raw
  a.update(es=es,ef=es+a['dur'],wait=es-raw);by[a['id']]=a
  a['start']=START+timedelta(days=es);a['finish']=START+timedelta(days=es+a['dur']-1 if a['dur'] else es)
 # float intentionally omitted: ordinary CPM float would misrepresent seasonal discontinuities.
 return tasks
BASE=calculate(BASE);REC=calculate(REC)
print('SCENARIOS')
for label,t in [('Two-crew case',BASE),('Recovery option',REC)]:
 print(label,[(a['id'],str(a['finish'])) for a in t if a['id'] in ['M100','O100','M200']])
# Validation: predecessors, calendar, and exclusive delivery resources.
for tasks in [BASE,REC]:
 by={a['id']:a for a in tasks}
 assert len(by)==len(tasks)
 for a in tasks:
  assert all(a['es']>=by[p]['ef'] for p in a['preds'])
  assert not a['off'] or seasonal_start(a['es'],a['dur'])==a['es']
 for resource in ['Fabrication line 1','Fabrication line 2','Heavy-lift spread','Strengthening crew A','Strengthening crew B','Strengthening crew C (additional)','Cargo spread']:
  series=sorted([a for a in tasks if a['resource']==resource],key=lambda a:a['es'])
  assert all(b['es']>=a['ef'] for a,b in zip(series,series[1:])),resource
wb=Workbook();wb.remove(wb.active)
navy='16324F';teal='087E8B';orange='FFF2CC';red='FCE4D6';green='E2F0D9'
def sheet(name,headers,rows,widths=None):
 s=wb.create_sheet(name);s.append(headers)
 for r in rows:s.append(r)
 s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
 for c in s[1]:c.fill=PatternFill('solid',fgColor=navy);c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 s.row_dimensions[1].height=32
 for row in s.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(c.value,date):c.number_format='dd-mmm-yyyy'
  if row[0].row%2==0:
   for c in row:c.fill=PatternFill('solid',fgColor='F2F6FA')
 for i in range(1,s.max_column+1):s.column_dimensions[get_column_letter(i)].width=widths[i-1] if widths and i<=len(widths) else 22
 s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=0;s.print_title_rows='1:1'
 return s
basis=[
('Status','PROVISIONAL MANAGEMENT SCHEDULE — Rev 0, 25 October 2026. Not ONGC approved; no actual progress inferred.'),
('Start assumption','1 November 2026 is a scenario from uploaded plans, NOT a confirmed award date. All relative days use this date as Day 1.'),
('Contract deadline','12 months from NOA under reviewed GCC 6.3.1.1 and SCC 2.0. For this scenario, use 31 October 2027 as conservative last workday; confirm legal date convention and actual NOA.'),
('Calendar','Onshore durations are elapsed calendar days assuming adequate shifts and staffing; not 5-day working durations. Validate holidays and productivity. Offshore tasks conservatively avoid 16 May–15 October inclusive.'),
('Weather treatment','An offshore activity that cannot fit wholly before 16 May is deferred wholly to 16 October; tasks are not split. This is a conservative planning rule, not a contractual automatic extension.'),
('Contract relief','SCC provides conditional monsoon time relief for inability to work due to adverse weather; associated extra costs are not payable under that clause. Obtain ONGC agreement; no extension is assumed in deadline comparison.'),
('Durations','All proposed durations are planner estimates, not measured quantity/productivity results or supplier commitments. Survey, approvals and procurement lead times require validation.'),
('Logic','Finish-to-start, zero lag. Successor starts next day after predecessor working finish. Milestone date is next-day boundary. Offshore seasonal waits are separately visible.'),
('Two-crew case','2 fabrication lines + 1 dedicated strengthening cell; 2 offshore strengthening crews; 1 sequential heavy-lift spread and 1 cargo spread. Candidate capacity only.'),
('Recovery option','Same fabrication lines and heavy-lift spread; first approved strengthening kits released after 20 rather than 35 days of cell production; ADDITIONAL crew C for NQO-NQD. Subject to package-specific release and availability.'),
('Support capacity','Support spread / transfers / accommodation must support simultaneous crews and lifting. SCC minimum 60 special personnel is a floor, not proof capacity is adequate. Verify POB, vessels, access and operational interfaces.'),
('Resource check','Exclusive fabrication-line, cargo, heavy-lift and strengthening-crew sequencing checked. This is NOT full quantity-loaded, labour-levelled P6 analysis. Survey/design teams and support capacities require validation.'),
('Weather allowance','10 reserved offshore days after planned offshore works, then 3 demobilisation days. This allowance is unvalidated; test against vessel weather limits and metocean history.'),
('Scope','Ten bridges; ICG/ICP/BHS jacket strengthening included; BHS new bridge/service transfer precedes old-bridge demolition. Other demolition allowance is provisional pending bridge allocation and approved methods.'),
('Quantities','Structural summary 1,142 MT vs bridge sum 1,141 MT unresolved. Demolition 357 MT allocation inconsistent across uploaded controls. No correction silently assumed.'),
('Completion','Internal completion includes all offshore work, specified punch clearance, dossiers and acceptance processing. ONGC approval durations are estimates; planned certificate date is not an acceptance commitment.'),
('Warranty','12-month warranty after actual completion / handover; shown separately and excluded from EPC duration. Final warranty security release depends on contract terms.'),
('Use','Recovery option is the proposed management target; two-crew case exposes seasonal risk. Neither is an approved baseline. Recalculate after actual NOA, supplier lead times, weather study and resource commitments.'),
('Source priority','Executed contract and controlling addenda govern. This workbook supersedes no source document and is a newly proposed overall schedule.'),
('Source references','GCC 2.1.1 / 6.3.1.1; Technical SCC 2.0 / 3.0; SOW Section 1 (including 1.3); Addendum_Structure_SOW; Stage28 / Level3 / 360-day workbook caveats; Stage34 and Stage35 planning assumptions.')]
sheet('READ ME',['Topic','Basis / requirement'],basis,[29,125])
headers=['ID','Phase','Package','Activity / milestone','Duration (calendar days)','FS predecessors','Planned start','Planned finish','Start day from NOA','Finish day from NOA','Seasonal wait days','Resource assumption','Accountable role','Release / acceptance condition','Status']
for label,tasks in [('Recovery Schedule',REC),('Two Crew Schedule',BASE)]:
 rows=[[a['id'],a['phase'],a['pkg'],a['name'],a['dur'],', '.join(a['preds']),a['start'],a['finish'],a['es']+1,a['ef'] if a['dur'] else a['ef']+1,a['wait'],a['resource'],a['owner'],a['gate'],'Proposed / unverified'] for a in tasks]
 s=sheet(label,headers,rows,[12,21,24,64,14,27,17,17,14,14,14,38,24,65,24])
 for r,a in enumerate(tasks,2):
  if a['dur']==0:
   for c in s[r]:c.fill=PatternFill('solid',fgColor=green);c.font=Font(bold=True)
  if a['wait']:s.cell(r,11).fill=PatternFill('solid',fgColor=red)
 with open(O/(label.replace(' ','_')+'.csv'),'w',newline='') as f:
  writer=csv.writer(f);writer.writerow(headers);writer.writerows(rows)
# Milestones / comparison.
bb={a['id']:a for a in BASE};rr={a['id']:a for a in REC}
msids=['M000','E030','P010','F010','F020','F030','O010','J010','J020','I010','I020','D020','I030','M100','O100','M200']
rows=[[i,rr[i]['name'],rr[i]['finish'],bb[i]['finish'],(bb[i]['finish']-rr[i]['finish']).days,'Conditional target'] for i in msids]
rows += [['DEADLINE','Unextended contractual last workday (scenario)',date(2027,10,31),date(2027,10,31),None,'Confirm actual NOA / contract convention'],['WARRANTY','Indicative warranty end: 12 months after actual handover',rr['M200']['finish'].replace(year=rr['M200']['finish'].year+1),None,None,'Rebase to actual acceptance date']]
sheet('Milestone Comparison',['ID','Milestone','Recovery option','Two-crew case','Recovery improvement days','Basis'],rows,[15,78,20,20,23,48])
# Weekly Gantt supports the whole contractual horizon plus base-case overrun.
weeks=[START+timedelta(days=7*i) for i in range(60)]
s=sheet('Weekly Gantt',['ID','Package / activity','Start','Finish']+weeks,[[a['id'],a['pkg']+' | '+a['name'],a['start'],a['finish']]+['' for _ in weeks] for a in REC],[12,75,17,17]+[4]*len(weeks));s.freeze_panes='E2';s.auto_filter.ref=f'A1:D{len(REC)+1}'
for j,w in enumerate(weeks,5):
 c=s.cell(1,j);c.value=w;c.number_format='dd-mmm';c.alignment=Alignment(textRotation=90);s.row_dimensions[1].height=55
 if any(date(2027,5,16)<=w+timedelta(days=d)<=date(2027,10,15) for d in range(7)):c.fill=PatternFill('solid',fgColor='A66F00')
for r,a in enumerate(REC,2):
 for j,w in enumerate(weeks,5):
  if a['start']<=w+timedelta(days=6) and a['finish']>=w:
   s.cell(r,j,'◆' if not a['dur'] else '■').fill=PatternFill('solid',fgColor=teal if a['off'] else '4F81BD')
   s.cell(r,j).font=Font(color='FFFFFF')
 s.row_dimensions[r].height=30
# Bridge closeout dates, differentiated from contractual package completion.
bridgeids={'ICW-ICG LD':'I010','ICG-ICP LD':'I030','BHS-SLQ':'D020','ICW-ICG UD':'RA01','ICG-ICP UD':'RA02','SLQ-WIS':'RA03','NQO-NQD':'RA04','WIN-NC':'RB01','SCA-SC1':'RB02','SCA-SCF':'RB03'}
sheet('Bridge Targets',['Bridge','Work type','Recovery offshore package target','Two-crew offshore target','Acceptance scope / caveat'],[[p,'Replacement' if i[0] in 'ID' else 'Refurbishment',rr[i]['finish'],bb[i]['finish'],'BHS includes old-bridge removal; other dates exclude later common residual demolition / final offshore punch, dossiers and overall acceptance'] for p,i in bridgeids.items()],[23,20,28,28,100])
risks=[
('R01','Actual NOA differs from 1 Nov','Recalculate the full network; date shift is not linear across monsoon.','Contracts','Before baseline','OPEN'),
('R02','Recovery requires third strengthening crew','Confirm competent crew C, supervision, accommodation and independent workfront.','Construction / Marine','Before commitment','OPEN'),
('R03','First-kit release after 20 days is unproven','Approve phased MTO, priority kits, inspection and transport; no work before QA release.','Engineering / Fabrication','Before production baseline','OPEN'),
('R04','Marine standby between campaigns','Cost reserved vessel days, accommodation, transfers and second-season remobilisation in base case.','Marine / Finance','Before charter','OPEN'),
('R05','Survey findings / demolition scope','Reconcile 1 MT; allocate 357 MT correctly; quantify other-bridge removal and engineered temporary works.','Engineering','Before procurement freeze','OPEN'),
('R06','Conditional monsoon relief','Submit compliant notices and obtain extension entitlement / dates; do not assume automatic relief or cost compensation.','Contracts','Before weather impact','OPEN'),
('R07','Support spread capacity','POB and simultaneous operations capacity may require multiple support vessels; minimum 60 does not establish adequacy.','Marine / ONGC','Before offshore release','OPEN'),
('R08','Durations and weather allowance not validated','Replace estimates with measured workpacks, rates, supplier delivery commitments and weather workability analysis.','Project Controls','Before baseline approval','OPEN'),
('R09','Approval turnaround / utility shutdown','Agree survey access, AFC review, service transfers, isolations, e-MB and acceptance windows.','Project Director / ONGC','Within first 30 days','OPEN'),
('R10','BHS removal sequence','New bridge and operational service transfer accepted before existing bridge demolition.','Technical Authority','Before demolition','HOLD'),
('R11','Shared engineering resources','Design team A requires structural, utility and temporary-works capacity concurrently; staffing numbers not established.','Engineering Manager','Before design start','OPEN'),
('R12','Unpriced project','Link dates to costs, charter terms, billing milestones and funding before commitment.','Finance / Commercial','Before sanction','OPEN')]
sheet('Constraints and Actions',['ID','Constraint','Required action','Owner','Decision point','Status'],risks,[12,48,105,30,28,16])
resource_rows=[('Fabrication line 1',1,'ICW-ICG LD then ICG-ICP LD','Serial logic enforced; no throughput proven'),('Fabrication line 2',1,'BHS-SLQ','Dedicated candidate bay'),('Strengthening cell',1,'All bridge and jacket kits','20-day first batch + 35-day full release is recovery assumption'),('Strengthening crews',3,'A: IC / SLQ; B: WIN / SCA; C: NQ','Recovery uses 3; comparator uses 2'),('Heavy-lift spread',1,'ICW LD → BHS install → BHS demolition → ICP LD → residual removal','Sequential major operations; technical suitability unverified'),('Cargo spread',1,'Sequential bridge transport','Voyage duration / deck suitability unverified'),('Support / accommodation spread','TBC','Survey, strengthening, heavy-lift and punch','Size for peak concurrent POB, logistics and ONGC operations'),('Engineering','2 discipline groups + specialist support','Parallel design and marine / utility studies','Headcount and review capacity TBC'),('QA / NDT / coating','TBC','Support concurrent fabrication lines, kits and offshore','Independent coverage and inspection hold points TBC')]
sheet('Resource Basis',['Resource','Planning quantity','Allocation','Qualification / limitation'],resource_rows,[30,22,86,95])
wb.save(O/'SHIVANSH_BRRP_Overall_Project_Schedule_Rev0.xlsx')
# concise printable management schedule (DOCX).
from docx import Document
from docx.shared import Inches,Pt
D=Document();D.add_heading('SHIVANSH / ONGC BRRP — Overall Project Schedule',0)
D.add_paragraph('Revision 0 | 25 October 2026 | Provisional management planning')
D.add_paragraph('Scenario commencement: 1 November 2026. Actual NOA / award is unconfirmed. Contract basis: 12 months from NOA, rather than a fixed 360-day limit. Dates and durations require approval and supplier / resource validation.')
D.add_heading('Recommended recovery target',1)
D.add_paragraph('Two fabrication lines and one strengthening cell; three offshore strengthening crews, one sequential heavy-lift spread and one cargo spread. Earlier first-kit release is required. These are capacity assumptions, not confirmed commitments.')
t=D.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['Activity / milestone','Start','Finish']):c.text=v
for i in ['S020','E030','E040','P010','F010','F020','F030','F040','O010','J010','J020','I010','I020','D020','I030','D030','O090','W010','O100','C010','C020','M200']:
 a=rr[i];cells=t.add_row().cells
 for c,v in zip(cells,[a['name'],a['start'].strftime('%d %b %Y'),a['finish'].strftime('%d %b %Y')]):c.text=v
D.add_heading('Seasonal and contractual assessment',1)
for label,ts in [('Recovery option',rr),('Two-crew comparison',bb)]:
 D.add_paragraph(f"{label}: offshore works complete {ts['M100']['finish']:%d %b %Y}; demobilisation {ts['O100']['finish']:%d %b %Y}; overall internal completion {ts['M200']['finish']:%d %b %Y}.")
D.add_paragraph('Conservative unextended contractual last workday for this scenario: 31 October 2027. Monsoon is excluded from planned offshore work, 16 May–15 October inclusive. Conditional contractual weather relief requires assessment and agreement; no extension or cost compensation is assumed.')
D.add_heading('Management release conditions',1)
for text in ['Confirm actual NOA, contract date convention and controlling addenda.','Validate activity durations from work quantities, approval periods, supplier lead times and weather workability.','Secure third strengthening crew and phased first-kit QA release; confirm supporting accommodation, access and logistics.','Reconcile structural 1,142 / 1,141 MT discrepancy and bridge-specific demolition quantities.','Obtain approved jacket strengthening, lift and BHS service-transfer / demolition methods before offshore execution.','Approve integrated cost, cash and charter plan. Maintain weekly updates and a six-week constraint look-ahead.']:
 D.add_paragraph(text,style='List Bullet')
D.add_paragraph('Sources reviewed: GCC 2.1.1 / 6.3.1.1; Technical SCC 2.0 / 3.0; SOW Section 1 / 1.3; structural addendum; uploaded Level3, Stage28 and 360-day planning files; Stage34/35 management controls. Full resource-loaded P6 validation remains outstanding. All dates are proposed, not ONGC commitments.')
D.save(O/'SHIVANSH_BRRP_Overall_Schedule_Management_Note_Rev0.docx')
Path('/workspace/project-review/final_schedule_data.json').write_text(json.dumps({'recovery':REC,'two_crew':BASE},default=str))
print('WORKBOOK AND NOTE SAVED',O)
